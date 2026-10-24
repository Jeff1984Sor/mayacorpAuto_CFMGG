"""Faturamento: contrato → caso → regras (fixo, recorrente, hora), despesas, Itens a faturar, NH/ND e NFS-e da fatura."""

from datetime import date

from app.modules.faturamento import service
from app.modules.faturamento.models import Frequencia, RegraCobranca, Termino, TipoRegra
from tests.test_api_fluxo import _login, cliente  # noqa: F401  (fixture)


def test_recorrencia_como_no_outlook():
    r = RegraCobranca(tipo=TipoRegra.recorrente, frequencia=Frequencia.mensal, intervalo=1, data_inicio=date(2026, 1, 31),
                      termino=Termino.ocorrencias, ocorrencias=3)
    assert [d for _, d in service.ocorrencias(r, date(2027, 1, 1))] == [date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 31)]
    r = RegraCobranca(tipo=TipoRegra.recorrente, frequencia=Frequencia.semanal, intervalo=2, data_inicio=date(2026, 1, 5),
                      termino=Termino.data, data_fim=date(2026, 2, 3))
    assert [d for _, d in service.ocorrencias(r, date(2027, 1, 1))] == [date(2026, 1, 5), date(2026, 1, 19), date(2026, 2, 2)]
    assert service.total_parcelas(r) == 3
    r = RegraCobranca(tipo=TipoRegra.recorrente, frequencia=Frequencia.anual, intervalo=1, data_inicio=date(2024, 2, 29),
                      termino=Termino.sem)
    assert [d for _, d in service.ocorrencias(r, date(2026, 12, 31))] == [date(2024, 2, 29), date(2025, 2, 28), date(2026, 2, 28)]


def _base(c, h, sufixo: str):
    cli = c.post("/api/nfse/tomadores", headers=h, json={
        "tipo": "PJ", "nome": f"Cliente Faturamento {sufixo}", "cidade": "São Paulo", "uf": "SP",
        "cpf_cnpj": {"A": "12345678000195", "B": "98765432000198"}.get(sufixo), "tipo_logradouro": "R",
        "logradouro": "RUA AUGUSTA", "numero": "100", "bairro": "CONSOLACAO", "cep": "01305000",
        "emails": [f"fin{sufixo}@cliente.com.br"]}).json()
    contrato = c.post("/api/faturamento/contratos", headers=h, json={"tomador_id": cli["id"], "nome": f"Consultivo {sufixo}"})
    assert contrato.status_code == 201, contrato.text
    contrato = contrato.json()
    caso = c.post(f"/api/faturamento/contratos/{contrato['id']}/casos", headers=h, json={"titulo": "Ação de cobrança"}).json()
    return cli, contrato, caso


def test_itens_a_faturar_e_faturas_nh_nd(cliente):  # noqa: F811
    h = _login(cliente)
    cli, contrato, caso = _base(cliente, h, "A")
    cid = contrato["id"]
    # mensal, 3 parcelas a partir de 10/01
    r = cliente.post(f"/api/faturamento/contratos/{cid}/regras", headers=h, json={
        "tipo": "recorrente", "descricao": "Mensalidade", "valor": "1000.00", "frequencia": "mensal",
        "data_inicio": "2026-01-10", "termino": "ocorrencias", "ocorrencias": 3})
    assert r.status_code == 201, r.text
    assert "3 parcelas" in r.json()["resumo"]
    assert cliente.post(f"/api/faturamento/contratos/{cid}/regras", headers=h, json={
        "tipo": "recorrente", "descricao": "x", "valor": "1", "frequencia": "mensal", "data_inicio": "2026-01-10",
        "termino": "ocorrencias"}).status_code == 422  # faltou a quantidade de parcelas
    cliente.post(f"/api/faturamento/contratos/{cid}/regras", headers=h, json={
        "tipo": "fixo", "descricao": "Êxito", "valor": "5000.00", "data_pagamento": "2026-02-15"})

    # por hora no caso: categoria → tabela → perfil do usuário
    cat = cliente.post("/api/config/cargos", headers=h, json={"nome": "Sócio"}).json()
    tab = cliente.post("/api/faturamento/tabelas", headers=h, json={
        "nome": "Tabela 2026", "itens": [{"cargo_id": cat["id"], "valor_hora": "800.00"}]}).json()
    eu = cliente.get("/api/auth/me", headers=h).json()
    assert cliente.patch(f"/api/config/usuarios/{eu["id"]}", headers=h, json={"cargo_id": cat["id"]}).status_code == 200
    cliente.post(f"/api/faturamento/casos/{caso['id']}/regras", headers=h, json={
        "tipo": "hora", "descricao": "Horas", "tabela_preco_id": tab["id"]})
    assert cliente.post(f"/api/faturamento/casos/{caso['id']}/timesheet", headers=h, json={
        "data": "2026-02-03", "minutos": 90, "descricao": "Petição inicial"}).status_code == 201
    desp = cliente.post(f"/api/faturamento/casos/{caso['id']}/despesas", headers=h, json={
        "data": "2026-02-04", "tipo": "custas", "descricao": "Guia de custas", "valor": "320.50"}).json()

    params = {"de": "2026-01-01", "ate": "2026-02-28", "cliente_ids": [cli["id"]]}
    r = cliente.get("/api/faturamento/itens-a-faturar", headers=h, params=params)
    assert r.status_code == 200, r.text
    dados = r.json()
    assert len(dados["clientes"]) == 1
    c0 = dados["clientes"][0]
    itens = c0["contratos"][0]["itens"]
    # 2 parcelas (jan, fev) + fixo + 1,5 h × 800 = 1200 (a regra do caso vence a do contrato)
    assert sorted(i["tipo"] for i in itens) == ["fixo", "hora", "recorrente", "recorrente"]
    assert c0["total"] == "8200.00"
    hora = next(i for i in itens if i["tipo"] == "hora")
    assert hora["valor"] == "1200.00" and hora["caso_numero"] == caso["numero"]

    # fatura só janeiro + horas
    escolhidas = [i["chave"] for i in itens if i["tipo"] == "hora" or i["data"] == "2026-01-10"]
    r = cliente.post("/api/faturamento/faturar", headers=h, json={"aba": "honorarios", **params, "chaves": escolhidas})
    assert r.status_code == 200, r.text
    nh = r.json()["faturas"][0]
    assert nh["codigo"].startswith("NH ") and nh["valor"] == "2200.00"
    restante = cliente.get("/api/faturamento/itens-a-faturar", headers=h, params=params).json()["clientes"][0]
    assert restante["total"] == "6000.00"
    # não cobra duas vezes
    r = cliente.post("/api/faturamento/faturar", headers=h, json={"aba": "honorarios", **params, "chaves": escolhidas})
    assert r.status_code == 400

    # despesas → ND (nunca NFS-e)
    r = cliente.get("/api/faturamento/itens-a-faturar", headers=h, params={**params, "aba": "despesas"}).json()
    item = r["clientes"][0]["contratos"][0]["itens"][0]
    assert item["chave"] == f"desp:{desp['id']}" and item["valor"] == "320.50"
    nd = cliente.post("/api/faturamento/faturar", headers=h, json={"aba": "despesas", **params, "chaves": [item["chave"]]}).json()
    nd = nd["faturas"][0]
    assert nd["codigo"].startswith("ND ")
    assert cliente.get(f"/api/nfse/faturas/{nd['id']}/pdf", headers=h).content[:5] == b"%PDF-"
    r = cliente.post("/api/faturamento/emitir-nfse", headers=h, json={"fatura_ids": [nd["id"]]})
    assert r.status_code == 400 and "Nota de Despesas" in r.json()["detail"]
    assert cliente.put(f"/api/faturamento/despesas/{desp['id']}", headers=h, json={
        "data": "2026-02-04", "descricao": "xy", "valor": "1"}).status_code == 400  # já faturada

    # lista de faturas mostra NH sem nota com o botão de emitir
    fat = cliente.get(f"/api/nfse/faturas/{nh['id']}", headers=h).json()
    assert fat["pode_emitir_nfse"] and fat["nfse_id"] is None and fat["contrato_id"] == cid
    assert cliente.get(f"/api/nfse/faturas/{nh['id']}/pdf", headers=h).content[:5] == b"%PDF-"

    # cancelar solta os itens
    r = cliente.post(f"/api/faturamento/faturas/{nh['id']}/cancelar", headers=h, json={"motivo": "valor errado"})
    assert r.status_code == 200, r.text
    assert cliente.get("/api/faturamento/itens-a-faturar", headers=h, params=params).json()["clientes"][0]["total"] == "8200.00"


def test_emitir_nfse_da_fatura_nh(cliente):  # noqa: F811
    h = _login(cliente)
    cli, contrato, _ = _base(cliente, h, "B")
    cliente.post(f"/api/faturamento/contratos/{contrato['id']}/regras", headers=h, json={
        "tipo": "fixo", "descricao": "Parecer", "valor": "3000.00", "data_pagamento": date.today().isoformat()})
    params = {"de": date.today().isoformat(), "ate": date.today().isoformat(), "cliente_ids": [cli["id"]]}
    itens = cliente.get("/api/faturamento/itens-a-faturar", headers=h, params=params).json()["clientes"][0]["contratos"][0]["itens"]
    nh = cliente.post("/api/faturamento/faturar", headers=h, json={
        "aba": "honorarios", **params, "chaves": [i["chave"] for i in itens]}).json()["faturas"][0]
    r = cliente.post("/api/faturamento/emitir-nfse", headers=h, json={"fatura_ids": [nh["id"]]})
    assert r.status_code == 200, r.text
    lote_id = r.json()["lote_id"]
    lote = cliente.get(f"/api/nfse/importacoes/{lote_id}", headers=h).json()
    assert lote["status"] in ("validando", "pronto"), lote
    r = cliente.post(f"/api/nfse/importacoes/{lote_id}/emitir", headers=h)
    assert r.status_code == 200, r.text
    fat = cliente.get(f"/api/nfse/faturas/{nh['id']}", headers=h).json()
    assert fat["nfse_id"] and fat["codigo"] == nh["codigo"] and not fat["pode_emitir_nfse"]
    # e não criou uma segunda fatura para a mesma nota
    assert cliente.get("/api/nfse/faturas", headers=h, params={"lote_id": lote_id}).json()["total"] == 1


def test_revisao_da_fatura(cliente):  # noqa: F811
    h = _login(cliente)
    cli, contrato, caso = _base(cliente, h, "C")
    cid = contrato["id"]
    cliente.post(f"/api/faturamento/contratos/{cid}/regras", headers=h, json={
        "tipo": "recorrente", "descricao": "Mensalidade", "valor": "2000.00", "frequencia": "mensal",
        "data_inicio": "2026-03-01", "termino": "sem"})
    tab = cliente.post("/api/faturamento/tabelas", headers=h, json={"nome": "Tabela C", "itens": []}).json()
    cliente.post(f"/api/faturamento/contratos/{cid}/regras", headers=h, json={
        "tipo": "hora", "descricao": "Horas", "tabela_preco_id": tab["id"]})
    params = {"de": "2026-03-01", "ate": "2026-04-30", "cliente_ids": [cli["id"]]}
    itens = cliente.get("/api/faturamento/itens-a-faturar", headers=h, params=params).json()["clientes"][0]["contratos"][0]["itens"]
    assert len(itens) == 2
    r = cliente.post("/api/faturamento/revisoes", headers=h, json={"aba": "honorarios", **params, "chaves": [itens[0]["chave"]]})
    assert r.status_code == 201, r.text
    rid = r.json()[0]["id"]
    # reservado: some de Itens a faturar
    restante = cliente.get("/api/faturamento/itens-a-faturar", headers=h, params=params).json()["clientes"][0]
    assert restante["qtd_itens"] == 1
    # fluxo mostra "Em revisão"
    fluxo = cliente.get("/api/faturamento/fluxo", headers=h).json()
    assert any(c["id"] == rid for c in fluxo["colunas"]["em_revisao"]["cards"])

    # incluir a parcela de abril, adicionar horas (valor/hora informado), lançamento avulso, excluir, editar
    disp = cliente.get(f"/api/faturamento/revisoes/{rid}/disponiveis", headers=h).json()
    assert any(d["chave"] == itens[1]["chave"] for d in disp)
    rev = cliente.post(f"/api/faturamento/revisoes/{rid}/incluir", headers=h, json={"chaves": [itens[1]["chave"]]}).json()
    assert rev["total"] == "4000.00"
    r = cliente.post(f"/api/faturamento/revisoes/{rid}/horas", headers=h, json={
        "caso_id": caso["id"], "data": "2026-04-10", "minutos": 120, "descricao": "reuniao", "valor_hora": "500"})
    assert r.status_code == 201, r.text
    rev = r.json()
    hora = next(i for i in rev["itens"] if i["tipo"] == "hora")
    assert hora["valor"] == "1000.00"
    rev = cliente.post(f"/api/faturamento/revisoes/{rid}/lancamentos", headers=h, json={
        "data": "2026-04-15", "descricao": "Ajuste acordado", "valor": "150.00"}).json()
    assert rev["total"] == "5150.00"
    rev = cliente.patch(f"/api/faturamento/revisoes/{rid}/itens/{hora['id']}", headers=h, json={
        "descricao": "Reunião com o cliente sobre a estratégia processual", "quantidade": "1.5"}).json()
    assert next(i for i in rev["itens"] if i["id"] == hora["id"])["valor"] == "750.00"
    parcela_abril = next(i for i in rev["itens"] if i["chave"] == itens[1]["chave"])
    rev = cliente.delete(f"/api/faturamento/revisoes/{rid}/itens/{parcela_abril['id']}", headers=h).json()
    assert rev["total"] == "2900.00"  # 2000 + 750 + 150
    # a parcela excluída volta para Itens a faturar
    assert cliente.get("/api/faturamento/itens-a-faturar", headers=h, params=params).json()["clientes"][0]["qtd_itens"] == 1
    ts = cliente.get(f"/api/faturamento/casos/{caso['id']}/timesheet", headers=h).json()
    assert ts[0]["minutos"] == 90 and ts[0]["descricao"].startswith("Reunião")

    # IA sem provedor configurado: erro claro
    r = cliente.post(f"/api/faturamento/revisoes/{rid}/ia", headers=h)
    assert r.status_code in (200, 400)

    r = cliente.post(f"/api/faturamento/revisoes/{rid}/aprovar", headers=h)
    assert r.status_code == 200, r.text
    rev = r.json()
    assert rev["status"] == "faturada" and rev["fatura_codigo"].startswith("NH ")
    fat = cliente.get(f"/api/nfse/faturas/{rev['fatura_id']}", headers=h).json()
    assert fat["valor"] == "2900.00" and fat["pode_emitir_nfse"]
    assert cliente.post(f"/api/faturamento/revisoes/{rid}/aprovar", headers=h).status_code == 400
    fluxo = cliente.get("/api/faturamento/fluxo", headers=h).json()
    assert any(c["id"] == fat["id"] for c in fluxo["colunas"]["aguardando_nfse"]["cards"])

    # descartar devolve os itens
    rid2 = cliente.post("/api/faturamento/revisoes", headers=h, json={
        "aba": "honorarios", **params, "chaves": [itens[1]["chave"]]}).json()[0]["id"]
    assert cliente.post(f"/api/faturamento/revisoes/{rid2}/descartar", headers=h, json={}).status_code == 200
    assert cliente.get("/api/faturamento/itens-a-faturar", headers=h, params=params).json()["clientes"][0]["qtd_itens"] == 1


def test_funcionario_sem_acesso_lanca_horas_mas_nao_entra(cliente):  # noqa: F811
    h = _login(cliente)
    cargo = cliente.post("/api/config/cargos", headers=h, json={"nome": "Estagiário"}).json()
    r = cliente.post("/api/config/usuarios", headers=h, json={
        "nome": "Ana Estagiária", "email": "ana.estagio@cfme.com.br", "perfil": "visualizador", "acesso_sistema": False,
        "cargo_id": cargo["id"], "cpf": "529.982.247-25", "oab": None, "cep": "01310-100", "logradouro": "Av. Paulista",
        "numero_endereco": "1000", "bairro": "Bela Vista", "cidade": "São Paulo", "uf": "SP"})
    assert r.status_code == 200, r.text
    ana = r.json()
    assert ana["acesso_sistema"] is False and ana["cargo_nome"] == "Estagiário" and ana["cep"] == "01310100"
    assert not ana["convite_pendente"]
    assert cliente.post("/api/auth/login", json={"email": "ana.estagio@cfme.com.br", "senha": "Cfme@2026"}).status_code == 401
    profs = cliente.get("/api/faturamento/profissionais", headers=h).json()
    assert any(p["id"] == ana["id"] and p["cargo_nome"] == "Estagiário" for p in profs)
    _, _, caso = _base(cliente, h, "D")
    r = cliente.post(f"/api/faturamento/casos/{caso['id']}/timesheet", headers=h, json={
        "profissional_id": ana["id"], "data": "2026-05-02", "minutos": 30, "descricao": "Pesquisa de jurisprudência"})
    assert r.status_code == 201 and r.json()["cargo_nome"] == "Estagiário"


def test_excluir_caso(cliente):  # noqa: F811
    h = _login(cliente)
    _, contrato, caso = _base(cliente, h, "E")
    vazio = cliente.post(f"/api/faturamento/contratos/{contrato['id']}/casos", headers=h, json={"titulo": "Consulta avulsa"}).json()
    assert cliente.delete(f"/api/faturamento/casos/{vazio['id']}", headers=h).json()["mensagem"] == "Caso excluído."
    cliente.post(f"/api/faturamento/casos/{caso['id']}/despesas", headers=h, json={
        "data": "2026-06-01", "tipo": "copias", "descricao": "Cópias", "valor": "10"})
    r = cliente.delete(f"/api/faturamento/casos/{caso['id']}", headers=h).json()
    assert "encerrado" in r["mensagem"]
    assert cliente.get(f"/api/faturamento/casos/{caso['id']}", headers=h).json()["status"] == "encerrado"


def test_ficha_do_caso_areas_campos_anexos_andamentos(cliente):  # noqa: F811
    h = _login(cliente)
    _, contrato, caso = _base(cliente, h, "F")
    area = cliente.post("/api/faturamento/areas", headers=h, json={"nome": "Trabalhista", "cor": "#1A8F5E"}).json()
    assert cliente.post("/api/faturamento/areas", headers=h, json={"nome": "trabalhista"}).status_code == 409
    fase = cliente.post("/api/faturamento/campos", headers=h, json={
        "nome": "Fase", "tipo": "lista", "opcoes": ["Conhecimento", "Recursal", "Execução"], "obrigatorio": True}).json()
    sindicato = cliente.post("/api/faturamento/campos", headers=h, json={"nome": "Sindicato", "tipo": "texto",
                                                                         "area_id": area["id"]}).json()
    assert cliente.post("/api/faturamento/campos", headers=h, json={"nome": "Lista vazia", "tipo": "lista"}).status_code == 422

    base = {"titulo": "Reclamação trabalhista", "area_id": area["id"], "polo": "passivo", "tribunal": "TRT-2",
            "vara": "12ª Vara do Trabalho", "comarca": "São Paulo", "valor_causa": "150000.00", "data_distribuicao": "2026-03-10"}
    r = cliente.put(f"/api/faturamento/casos/{caso['id']}", headers=h, json={**base, "campos": {}})
    assert r.status_code == 400 and "Fase" in r.json()["detail"]  # obrigatório
    r = cliente.put(f"/api/faturamento/casos/{caso['id']}", headers=h, json={**base, "campos": {str(fase["id"]): "Outra"}})
    assert r.status_code == 400  # fora das opções
    r = cliente.put(f"/api/faturamento/casos/{caso['id']}", headers=h, json={
        **base, "campos": {str(fase["id"]): "Recursal", str(sindicato["id"]): "Sindicato X", "999": "ignorado"}})
    assert r.status_code == 200, r.text
    c = r.json()
    assert c["area_nome"] == "Trabalhista" and c["polo"] == "passivo" and c["valor_causa"] == "150000.00"
    assert c["campos"] == {str(fase["id"]): "Recursal", str(sindicato["id"]): "Sindicato X"}
    assert any(a["id"] == area["id"] and a["qtd_casos"] == 1 for a in cliente.get("/api/faturamento/areas", headers=h).json())

    # anexos
    r = cliente.post(f"/api/faturamento/casos/{caso['id']}/anexos", headers=h, data={"descricao": "Inicial", "categoria": "Petição"},
                     files=[("arquivos", ("inicial.pdf", b"%PDF-1.4 teste", "application/pdf")),
                            ("arquivos", ("procuracao.docx", b"PK..", "application/octet-stream"))])
    assert r.status_code == 201, r.text
    anexos = r.json()
    assert len(anexos) == 2 and anexos[0]["categoria"] == "Petição"
    assert cliente.get(f"/api/faturamento/anexos/{anexos[0]['id']}", headers=h).content == b"%PDF-1.4 teste"
    assert cliente.post(f"/api/faturamento/casos/{caso['id']}/anexos", headers=h,
                        files=[("arquivos", ("virus.exe", b"MZ", "application/octet-stream"))]).status_code == 400
    assert cliente.delete(f"/api/faturamento/anexos/{anexos[1]['id']}", headers=h).status_code == 200
    assert cliente.get(f"/api/faturamento/casos/{caso['id']}", headers=h).json()["qtd_anexos"] == 1

    # andamentos
    r = cliente.post(f"/api/faturamento/casos/{caso['id']}/andamentos", headers=h, json={
        "data": "2026-04-01", "tipo": "audiencia", "titulo": "Contestação", "prazo": "2026-04-15"})
    assert r.status_code == 201, r.text
    and1 = r.json()
    assert and1["atrasado"] is True  # prazo no passado e não concluído
    r = cliente.put(f"/api/faturamento/andamentos/{and1['id']}", headers=h, json={**and1, "concluido": True})
    assert r.json()["atrasado"] is False
    cliente.post(f"/api/faturamento/casos/{caso['id']}/andamentos", headers=h, json={
        "data": "2026-05-01", "tipo": "audiencia", "titulo": "Audiência inicial"})
    lista = cliente.get(f"/api/faturamento/casos/{caso['id']}/andamentos", headers=h).json()
    assert [a["titulo"] for a in lista] == ["Audiência inicial", "Contestação"]
    assert cliente.delete(f"/api/faturamento/andamentos/{and1['id']}", headers=h).status_code == 200

    # IA sem provedor configurado: mensagem clara
    assert cliente.post(f"/api/faturamento/casos/{caso['id']}/ia/resumo", headers=h).status_code == 400
    assert cliente.post("/api/faturamento/casos/ia/preencher", headers=h, data={"texto": "Processo 0001234-56.2026.5.02.0012"}).status_code == 400
    # campo usado só desativa
    assert "desativado" in cliente.delete(f"/api/faturamento/campos/{fase['id']}", headers=h).json()["mensagem"]


def test_dashboard_mostra_todas_as_faturas(cliente):  # noqa: F811
    h = _login(cliente)
    d = cliente.get("/api/dashboard/faturamento", headers=h).json()
    tipos = {x["tipo"] for x in d["notas"]}
    assert "ND" in tipos and "NH" in tipos  # criadas nos testes anteriores
    nd = next(x for x in d["notas"] if x["tipo"] == "ND")
    assert nd["nf"] is None and nd["nfse"].startswith("Não se aplica") and nd["origem"] == "Contrato"
    assert len({x["chave"] for x in d["notas"]}) == len(d["notas"])


def test_exportar_importar_contratos_e_casos(cliente):  # noqa: F811
    import io

    from openpyxl import load_workbook

    h = _login(cliente)
    cli, contrato, caso = _base(cliente, h, "G")
    r = cliente.put(f"/api/faturamento/casos/{caso['id']}", headers=h, json={"titulo": "Caso G", "tipo": "judicial_federal"})
    assert r.status_code == 200 and r.json()["tipo_nome"] == "Judicial Federal"
    assert any(c["id"] == caso["id"] for c in cliente.get("/api/faturamento/casos", headers=h, params={"tipo": "judicial_federal"}).json())

    # contratos: exporta, edita e importa de volta + linha nova
    x = cliente.get("/api/faturamento/contratos/exportar.xlsx", headers=h)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    wb = load_workbook(io.BytesIO(x.content))
    ws = wb["Contratos"]
    linha = next(r for r in ws.iter_rows(min_row=2) if r[0].value == contrato["numero"])
    linha[3].value = "Consultivo G (renomeado)"
    ws.append([None, cli["nome"], None, "Contrato importado", "Objeto", "01/02/2026", None, None, None, "Sim", None])
    ws.append([None, "Cliente que não existe", None, "X", None, None, None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    r = cliente.post("/api/faturamento/contratos/importar", headers=h, files={"arquivo": ("c.xlsx", buf.getvalue())}).json()
    assert r["criados"] == 1 and r["atualizados"] >= 1 and r["qtd_erros"] == 1
    assert cliente.get(f"/api/faturamento/contratos/{contrato['id']}", headers=h).json()["nome"] == "Consultivo G (renomeado)"

    # casos: modelo + importação com tipo, área nova e campo adicional
    m = load_workbook(io.BytesIO(cliente.get("/api/faturamento/casos/modelo.xlsx", headers=h).content))
    ws = m["Casos"]
    cab = [c.value for c in ws[1]]
    assert cab[:4] == ["Nº", "Nº do contrato", "Título", "Tipo do caso"]
    ws.delete_rows(2)
    nova = {k: None for k in cab}
    nova.update({"Nº do contrato": contrato["numero"], "Título": "Caso importado", "Tipo do caso": "Judicial Trabalhista",
                 "Área": "Previdenciário", "Polo": "Réu", "Valor da causa": "10.000,50", "Distribuição": "05/03/2026"})
    ws.append([nova[k] for k in cab])
    ws.append([None, 99999, "Sem contrato"] + [None] * (len(cab) - 3))
    buf = io.BytesIO()
    m.save(buf)
    r = cliente.post("/api/faturamento/casos/importar", headers=h, files={"arquivo": ("k.xlsx", buf.getvalue())}).json()
    assert r["criados"] == 1 and r["qtd_erros"] == 1, r
    novo = next(c for c in cliente.get("/api/faturamento/casos", headers=h, params={"q": "Caso importado"}).json())
    assert novo["tipo"] == "judicial_trabalhista" and novo["area_nome"] == "Previdenciário" and novo["polo"] == "passivo"
    assert novo["valor_causa"] == "10000.50" and novo["data_distribuicao"] == "2026-03-05"
    assert cliente.get("/api/faturamento/casos/exportar.xlsx", headers=h).content[:2] == b"PK"


def test_comprovante_da_despesa_vai_para_anexos(cliente):  # noqa: F811
    h = _login(cliente)
    _, _, caso = _base(cliente, h, "H")
    d = cliente.post(f"/api/faturamento/casos/{caso['id']}/despesas", headers=h, json={
        "data": "2026-10-01", "tipo": "deslocamento", "descricao": "Uber fórum", "valor": "30"}).json()
    r = cliente.post(f"/api/faturamento/despesas/{d['id']}/comprovante", headers=h, files={"arquivo": ("uber.pdf", b"%PDF-1", "application/pdf")})
    assert r.status_code == 200 and r.json()["tem_comprovante"]
    anexos = cliente.get(f"/api/faturamento/casos/{caso['id']}/anexos", headers=h).json()
    assert len(anexos) == 1 and anexos[0]["categoria"] == "Despesas" and anexos[0]["despesa_id"] == d["id"]
    assert cliente.delete(f"/api/faturamento/anexos/{anexos[0]['id']}", headers=h).status_code == 400
    assert cliente.delete(f"/api/faturamento/despesas/{d['id']}", headers=h).status_code == 200
    assert cliente.get(f"/api/faturamento/casos/{caso['id']}/anexos", headers=h).json() == []


def test_judiciario_tribunais_comarcas_varas_instancia(cliente):  # noqa: F811
    h = _login(cliente)
    est = cliente.get("/api/faturamento/judiciario/tribunais", headers=h, params={"tipo": "judicial_estadual"}).json()
    assert len(est) == 27 and any(t["sigla"] == "TJSP" for t in est) and any(t["sigla"] == "TJDFT" for t in est)
    fed = cliente.get("/api/faturamento/judiciario/tribunais", headers=h, params={"tipo": "judicial_federal"}).json()
    assert [t["sigla"] for t in fed] == ["TRF1", "TRF2", "TRF3", "TRF4", "TRF5", "TRF6"]
    trab = cliente.get("/api/faturamento/judiciario/tribunais", headers=h, params={"tipo": "judicial_trabalhista"}).json()
    assert len(trab) == 24
    sup = cliente.get("/api/faturamento/judiciario/tribunais", headers=h, params={"instancia": "superior"}).json()
    assert {t["sigla"] for t in sup} >= {"STF", "STJ", "TST"}
    com = cliente.get("/api/faturamento/judiciario/comarcas", headers=h, params={"tribunal": "TJSP"}).json()
    assert "Campinas" in com and len(com) == 645
    assert "Vitória (ES)" in cliente.get("/api/faturamento/judiciario/comarcas", headers=h, params={"tribunal": "TRF2"}).json()
    _, _, caso = _base(cliente, h, "J")
    r = cliente.put(f"/api/faturamento/casos/{caso['id']}", headers=h, json={
        "titulo": "Caso J", "tipo": "judicial_estadual", "instancia": "primeira", "tribunal": "TJSP", "comarca": "Campinas",
        "vara": "3ª Vara Cível de Campinas"})
    assert r.status_code == 200 and r.json()["instancia_nome"] == "1ª instância"
    varas = cliente.get("/api/faturamento/judiciario/varas", headers=h, params={"tribunal": "TJSP", "comarca": "Campinas"}).json()
    assert varas[0] == "3ª Vara Cível de Campinas" and "1ª Vara Cível" in varas


def test_tarefas_crud_filtros_e_anexos(cliente):  # noqa: F811
    from datetime import timedelta

    h = _login(cliente)
    _, _, caso = _base(cliente, h, "T")
    hoje = date.today()
    assert cliente.post(f"/api/faturamento/casos/{caso['id']}/andamentos", headers=h, json={
        "data": hoje.isoformat(), "tipo": "prazo", "titulo": "x"}).status_code == 422  # prazo virou tarefa
    t1 = cliente.post("/api/faturamento/tarefas", headers=h, json={
        "caso_id": caso["id"], "titulo": "Protocolar contestação", "tipo": "prazo", "prioridade": "urgente",
        "prazo": (hoje - timedelta(days=1)).isoformat(), "etiquetas": ["contencioso", " contencioso "]})
    assert t1.status_code == 201, t1.text
    t1 = t1.json()
    assert t1["atrasada"] and t1["etiquetas"] == ["contencioso"] and t1["responsavel_id"] and t1["caso_numero"] == caso["numero"]
    avulsa = cliente.post("/api/faturamento/tarefas", headers=h, json={
        "titulo": "Renovar certificado digital", "tipo": "administrativo", "prazo": (hoje + timedelta(days=3)).isoformat(), "hora": "14:30"}).json()
    assert avulsa["caso_id"] is None and avulsa["hora"] == "14:30"
    assert cliente.post("/api/faturamento/tarefas", headers=h, json={"titulo": "x hora", "hora": "25:00"}).status_code == 422

    r = cliente.get("/api/faturamento/tarefas", headers=h, params={"situacao": "atrasadas"}).json()
    assert any(t["id"] == t1["id"] for t in r["itens"]) and r["resumo"]["atrasadas"] >= 1
    assert any(t["id"] == avulsa["id"] for t in cliente.get("/api/faturamento/tarefas", headers=h, params={"situacao": "semana"}).json()["itens"])
    assert [t["id"] for t in cliente.get("/api/faturamento/tarefas", headers=h, params={"caso_id": caso["id"]}).json()["itens"]] == [t1["id"]]
    assert cliente.get("/api/faturamento/tarefas", headers=h, params={"sem_caso": True, "q": "certificado"}).json()["itens"][0]["id"] == avulsa["id"]
    assert cliente.get(f"/api/faturamento/casos/{caso['id']}", headers=h).json()["qtd_tarefas_abertas"] == 1
    assert "contencioso" in cliente.get("/api/faturamento/tarefas/etiquetas", headers=h).json()

    # anexos: tarefa do caso aparece na pasta Tarefas do caso; tarefa avulsa não
    a = cliente.post(f"/api/faturamento/tarefas/{t1['id']}/anexos", headers=h, files=[("arquivos", ("minuta.docx", b"PK", "application/octet-stream"))])
    assert a.status_code == 201, a.text
    anexos_caso = cliente.get(f"/api/faturamento/casos/{caso['id']}/anexos", headers=h).json()
    assert anexos_caso[0]["categoria"] == "Tarefas" and anexos_caso[0]["tarefa_id"] == t1["id"]
    b = cliente.post(f"/api/faturamento/tarefas/{avulsa['id']}/anexos", headers=h, files=[("arquivos", ("boleto.pdf", b"%PDF", "application/pdf"))]).json()
    assert cliente.get(f"/api/faturamento/anexos/{b[0]['id']}", headers=h).content == b"%PDF"
    assert cliente.get(f"/api/faturamento/tarefas/{avulsa['id']}", headers=h).json()["qtd_anexos"] == 1

    # status e exclusão
    r = cliente.patch(f"/api/faturamento/tarefas/{t1['id']}/status", headers=h, json={"status": "concluida"}).json()
    assert r["status"] == "concluida" and r["concluida_em"] == hoje.isoformat() and not r["atrasada"]
    assert cliente.delete(f"/api/faturamento/tarefas/{avulsa['id']}", headers=h).status_code == 200
    assert cliente.get(f"/api/faturamento/anexos/{b[0]['id']}", headers=h).status_code == 404
