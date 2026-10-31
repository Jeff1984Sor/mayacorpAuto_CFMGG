"""Financeiro: contas bancárias, extrato, transferência, contas a pagar (rateio, recorrência, baixa), planilhas,
recebimento de fatura no extrato, fluxo, DRE, projeções e impostos."""

import io
from datetime import date, timedelta
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from tests.test_api_fluxo import _login, cliente  # noqa: F401  (fixture)

F = "/api/financeiro"


def _conta(c, h, nome, tipo="corrente", saldo="1000.00"):
    r = c.post(f"{F}/contas", headers=h, json={"nome": nome, "banco": "Itaú", "tipo": tipo, "saldo_inicial": saldo})
    assert r.status_code == 201, r.text
    return r.json()


def _cat(c, h, nome, tipo="despesa"):
    return next(x for x in c.get(f"{F}/categorias", headers=h).json() if x["nome"] == nome and x["tipo"] == tipo)


def test_contas_extrato_transferencia_e_contas_a_pagar(cliente):  # noqa: F811
    h = _login(cliente)
    cc = _conta(cliente, h, "Itaú PJ")
    inv = _conta(cliente, h, "Investimento CDB", "investimento", "50000.00")
    cats = cliente.get(f"{F}/categorias", headers=h).json()
    assert any(x["nome"] == "Honorários advocatícios" for x in cats)  # plano de contas padrão
    aluguel = _cat(cliente, h, "Aluguel e condomínio")
    tec = _cat(cliente, h, "Tecnologia e sistemas")
    sub = cliente.post(f"{F}/categorias/{tec['id']}/subcategorias", headers=h, json={"nome": "Licenças"}).json()

    hoje = date.today()
    r = cliente.post(f"{F}/contas/{cc['id']}/lancamentos", headers=h, json={
        "data": hoje.isoformat(), "descricao": "Tarifa", "tipo": "debito", "valor": "50.00"})
    assert r.status_code == 201
    r = cliente.post(f"{F}/transferencias", headers=h, json={"origem_id": inv["id"], "destino_id": cc["id"],
                                                             "data": hoje.isoformat(), "valor": "2000.00"})
    assert r.status_code == 201, r.text
    ext = cliente.get(f"{F}/contas/{cc['id']}/extrato", headers=h).json()
    assert ext["saldo_final"] == "2950.00" and ext["linhas"][-1]["saldo"] == "2950.00"
    contas = {c["id"]: c for c in cliente.get(f"{F}/contas", headers=h).json()}
    assert contas[inv["id"]]["saldo_atual"] == "48000.00"
    # excluir uma perna apaga o par
    transf = next(x for x in ext["linhas"] if x["origem"] == "transferencia")
    assert cliente.delete(f"{F}/lancamentos/{transf['id']}", headers=h).status_code == 200
    contas = {c["id"]: c for c in cliente.get(f"{F}/contas", headers=h).json()}
    assert contas[inv["id"]]["saldo_atual"] == "50000.00" and contas[cc["id"]]["saldo_atual"] == "950.00"

    # fornecedor + conta com rateio e recorrência mensal (3 parcelas)
    forn = cliente.post(f"{F}/fornecedores", headers=h, json={"nome": "Imobiliária Centro", "tipo": "PJ",
                                                              "cpf_cnpj": "12.345.678/0001-95"}).json()
    r = cliente.post(f"{F}/contas-pagar", headers=h, json={
        "fornecedor_id": forn["id"], "data_vencimento": (hoje - timedelta(days=2)).isoformat(), "forma_pagamento": "boleto",
        "rateio": [{"categoria_id": aluguel["id"], "valor": "700.00"},
                   {"categoria_id": tec["id"], "subcategoria_id": sub["id"], "valor": "300.00"}],
        "recorrencia_meses": 1, "repeticoes": 3})
    assert r.status_code == 201, r.text
    serie = r.json()
    assert len(serie) == 3 and serie[0]["valor"] == "1000.00" and serie[0]["descricao"] == "Imobiliária Centro"
    assert serie[0]["rateio"][0]["percentual"] == "70.00" and serie[0]["status"] == "atrasado"
    assert serie[2]["parcela"] == 3 and serie[2]["parcelas_total"] == 3

    lista = cliente.get(f"{F}/contas-pagar", headers=h, params={"fornecedor_id": forn["id"]}).json()
    assert lista["kpis"]["qtd_atrasado"] == 1 and lista["kpis"]["em_aberto"] == "3000.00"
    # editar a série muda as demais em aberto, mantendo o vencimento
    r = cliente.put(f"{F}/contas-pagar/{serie[1]['id']}", headers=h, json={
        "fornecedor_id": forn["id"], "descricao": "Aluguel sala 12", "data_vencimento": serie[1]["data_vencimento"],
        "rateio": [{"categoria_id": aluguel["id"], "valor": "1100.00"}], "escopo": "serie"})
    assert r.status_code == 200, r.text
    c3 = cliente.get(f"{F}/contas-pagar/{serie[2]['id']}", headers=h).json()
    assert c3["valor"] == "1100.00" and c3["descricao"] == "Aluguel sala 12" and c3["data_vencimento"] == serie[2]["data_vencimento"]

    # baixa: débito no extrato; reabrir devolve o saldo
    r = cliente.post(f"{F}/contas-pagar/{serie[0]['id']}/pagar", headers=h, json={
        "data": hoje.isoformat(), "conta_bancaria_id": cc["id"], "juros": "10.00"})
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "pago" and r.json()["valor_pago"] == "1010.00"
    contas = {c["id"]: c for c in cliente.get(f"{F}/contas", headers=h).json()}
    assert contas[cc["id"]]["saldo_atual"] == "-60.00"
    deb = cliente.get(f"{F}/contas/{cc['id']}/extrato", headers=h).json()["linhas"][-1]
    assert deb["origem"] == "conta_pagar" and not deb["editavel"]
    assert cliente.delete(f"{F}/lancamentos/{deb['id']}", headers=h).status_code == 400
    assert cliente.put(f"{F}/contas-pagar/{serie[0]['id']}", headers=h, json={
        "descricao": "x", "data_vencimento": hoje.isoformat(), "rateio": [{"categoria_id": aluguel["id"], "valor": "1"}]}).status_code == 400
    assert cliente.post(f"{F}/contas-pagar/{serie[0]['id']}/reabrir", headers=h).status_code == 200
    contas = {c["id"]: c for c in cliente.get(f"{F}/contas", headers=h).json()}
    assert contas[cc["id"]]["saldo_atual"] == "950.00"
    # categoria em uso não é excluída
    assert cliente.delete(f"{F}/categorias/{aluguel['id']}", headers=h).status_code == 400
    # excluir a série a partir da 2ª
    r = cliente.delete(f"{F}/contas-pagar/{serie[1]['id']}", headers=h, params={"escopo": "serie"})
    assert r.status_code == 200 and "2 contas" in r.json()["mensagem"]

    # relatórios respondem e batem
    fluxo = cliente.get(f"{F}/fluxo-caixa", headers=h, params={"conta_ids": [cc["id"]]}).json()
    assert fluxo["saldo_final"] == 950
    dre = cliente.get(f"{F}/dre", headers=h, params={"mensal": True}).json()
    assert any(lin["chave"] == "resultado_liquido" for lin in dre["linhas"])
    proj = cliente.get(f"{F}/projecao-despesas", headers=h).json()
    assert len(proj["linhas"]) == 12 and proj["linhas"][0]["contas_a_pagar"] == 1000
    prev = cliente.get(f"{F}/previsao-caixa", headers=h, params={"horizonte": 30, "saldo_minimo": 5000}).json()
    assert prev["kpis"]["investimento"] == 50000 and prev["plano_transferencias"]  # corrente fica abaixo do mínimo
    assert prev["alerta"]["nivel"] == "atencao"
    for url in ("/fluxo-caixa.xlsx", "/dre.xlsx", "/contas-pagar/exportar.xlsx", f"/contas/{cc['id']}/extrato.xlsx"):
        r = cliente.get(F + url, headers=h)
        assert r.status_code == 200 and r.content[:2] == b"PK", url


def test_planilhas_de_importacao(cliente):  # noqa: F811
    h = _login(cliente)
    modelo = cliente.get(f"{F}/contas-pagar/modelo.xlsx", headers=h)
    wb = load_workbook(io.BytesIO(modelo.content))
    assert wb.sheetnames == ["Contas", "Instruções"] and wb["Contas"]["A1"].value.startswith("Vencimento")
    ws = wb["Contas"]
    ws.append(["10/12/2026", None, "Marketing", "Eventos", "", "Gráfica Rápida", "1.234,56", "PIX", "NF 9", ""])
    ws.append(["data ruim", None, "Marketing", None, "x", None, "10", None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    r = cliente.post(f"{F}/contas-pagar/importar", headers=h, files={"arquivo": ("contas.xlsx", buf.getvalue())})
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["importadas"] == 3 and res["qtd_erros"] == 1 and "Marketing" in res["categorias_novas"]
    lista = cliente.get(f"{F}/contas-pagar", headers=h, params={"q": "Gráfica"}).json()
    assert lista["itens"][0]["valor"] == "1234.56" and lista["itens"][0]["forma_pagamento"] == "pix"

    conta = _conta(cliente, h, "Bradesco")
    wb = Workbook()
    ws = wb.active
    ws.append(["Data", "Histórico", "Valor"])
    ws.append([date.today(), "PIX recebido cliente", 500])
    ws.append([date.today(), "Tarifa pacote", "-35,00"])
    buf = io.BytesIO()
    wb.save(buf)
    r = cliente.post(f"{F}/contas/{conta['id']}/importar", headers=h, files={"arquivo": ("ext.xlsx", buf.getvalue())})
    assert r.status_code == 200, r.text
    assert r.json()["importados"] == 2
    ext = cliente.get(f"{F}/contas/{conta['id']}/extrato", headers=h).json()
    assert ext["entradas"] == "500.00" and ext["saidas"] == "35.00"


def test_modelo_e_importacao_de_fornecedores(cliente):  # noqa: F811
    h = _login(cliente)
    existente = cliente.post(f"{F}/fornecedores", headers=h, json={"nome": "Contabilidade Antiga", "tipo": "PJ",
                                                                   "cpf_cnpj": "45.997.418/0001-53", "telefone": "(11) 1111-1111"}).json()
    r = cliente.get(f"{F}/fornecedores/modelo.xlsx", headers=h)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Fornecedores", "Instruções"]
    ws = wb["Fornecedores"]
    assert [c.value for c in ws[1]][:3] == ["Nome", "Tipo (PF/PJ)", "CPF/CNPJ"]
    ws.delete_rows(2)
    ws.append(["Maria Tradutora", None, "529.982.247-25", "maria@exemplo.com", None, "maria@exemplo.com", None, None, None,
               "Serviços de tradução", None, "Sim"])
    ws.append(["Contabilidade Nova", "PJ", "45997418000153", "contab@exemplo.com", None, None, None, None, None, None, None, None])
    ws.append(["Sem documento certo", "PJ", "123", None, None, None, None, None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    r = cliente.post(f"{F}/fornecedores/importar", headers=h, files={"arquivo": ("forn.xlsx", buf.getvalue())})
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["criados"] == 1 and res["atualizados"] == 1 and res["qtd_erros"] == 1, res
    assert "Linha 4" in res["erros"][0] and "CNPJ" in res["erros"][0]
    assert "Serviços de tradução" in res["categorias_novas"]
    lista = {f["cpf_cnpj"]: f for f in cliente.get(f"{F}/fornecedores", headers=h).json()}
    maria = lista["52998224725"]
    assert maria["tipo"] == "PF" and maria["nome"] == "Maria Tradutora" and maria["categoria_padrao_id"]
    atual = lista["45997418000153"]
    assert atual["id"] == existente["id"] and atual["nome"] == "Contabilidade Nova"
    assert atual["email"] == "contab@exemplo.com" and atual["telefone"] == "(11) 1111-1111"  # célula vazia não apaga


def test_recebimento_de_fatura_vira_credito(cliente):  # noqa: F811
    h = _login(cliente)
    conta = _conta(cliente, h, "Conta Recebimentos", saldo="0")
    cfg = cliente.get(f"{F}/config", headers=h).json()
    cfg["conta_recebimento_id"] = conta["id"]
    assert cliente.put(f"{F}/config", headers=h, json=cfg).status_code == 200
    # fatura de despesas (ND) do faturamento
    cli = cliente.post("/api/nfse/tomadores", headers=h, json={"tipo": "PJ", "nome": "Cliente Financeiro"}).json()
    contrato = cliente.post("/api/faturamento/contratos", headers=h, json={"tomador_id": cli["id"], "nome": "Contencioso"}).json()
    caso = cliente.post(f"/api/faturamento/contratos/{contrato['id']}/casos", headers=h, json={"titulo": "Processo X"}).json()
    hoje = date.today().isoformat()
    cliente.post(f"/api/faturamento/casos/{caso['id']}/despesas", headers=h, json={
        "data": hoje, "tipo": "custas", "descricao": "Custas iniciais", "valor": "250.00"})
    params = {"de": hoje, "ate": hoje, "cliente_ids": [cli["id"]]}
    item = cliente.get("/api/faturamento/itens-a-faturar", headers=h, params={**params, "aba": "despesas"}).json()
    chave = item["clientes"][0]["contratos"][0]["itens"][0]["chave"]
    nd = cliente.post("/api/faturamento/faturar", headers=h, json={"aba": "despesas", **params, "chaves": [chave]}).json()["faturas"][0]
    r = cliente.post(f"/api/nfse/faturas/{nd['id']}/baixa", headers=h, json={"data": hoje})
    assert r.status_code == 200, r.text
    ext = cliente.get(f"{F}/contas/{conta['id']}/extrato", headers=h).json()
    cred = ext["linhas"][-1]
    assert cred["origem"] == "fatura" and cred["valor"] == "250.00" and cred["fatura_codigo"] == nd["codigo"]
    assert cred["categoria_nome"] == "Reembolso de despesas"
    assert cliente.delete(f"/api/nfse/faturas/{nd['id']}/baixa", headers=h).status_code == 200
    assert cliente.get(f"{F}/contas/{conta['id']}/extrato", headers=h).json()["linhas"] == []
    rel = cliente.get(f"{F}/relatorio-faturamento", headers=h, params={"tomador_id": cli["id"]}).json()
    assert rel["resumo"]["qtd"] == 1 and rel["resumo"]["total"] == 250 and rel["por_tipo"][0]["tipo"] == "ND"
    assert cliente.get(f"{F}/relatorio-faturamento.xlsx", headers=h).content[:2] == b"PK"


def test_projecao_de_impostos(cliente):  # noqa: F811
    h = _login(cliente)
    r = cliente.get(f"{F}/impostos", headers=h)
    assert r.status_code == 200, r.text
    p = r.json()
    assert len(p["linhas"]) == 25 and p["resumo"]["regime"] in ("simples", "presumido")
    eid = p["emissora_id"]
    # Simples Anexo IV com receita informada pelo contador
    cfg = {**p["config"], "regime": "simples", "anexo": "IV", "recarregar_tabela": True}
    cfg.pop("aliquota_iss_emissora")
    cfg.pop("regime_cadastro")
    p = cliente.put(f"{F}/impostos/{eid}/config", headers=h, json=cfg).json()
    assert p["resumo"]["regime"] == "simples" and len(p["config"]["faixas"]) == 6
    mes_passado = (date.today().replace(day=1) - timedelta(days=1)).strftime("%Y-%m")
    p = cliente.post(f"{F}/impostos/{eid}/receitas-informadas", headers=h, json={"mes": mes_passado, "valor": "100000"}).json()
    linha = next(x for x in p["linhas"] if x["mes"] == mes_passado)
    assert linha["origem"] == "contador" and "DAS" in linha["impostos"]
    s = cliente.post(f"{F}/impostos/{eid}/simular", headers=h, json={"valor": "10000", "mes": date.today().isoformat()}).json()
    assert float(s["imposto"]) > 0 and float(s["liquido"]) < 10000
    # Presumido: IRPJ/CSLL só no fechamento do trimestre
    cfg["regime"] = "presumido"
    p = cliente.put(f"{F}/impostos/{eid}/config", headers=h, json=cfg).json()
    tri = [x for x in p["linhas"] if x["mes"].endswith(("-03", "-06", "-09", "-12"))]
    assert all("IRPJ" in x["impostos"] for x in tri)
    assert all("IRPJ" not in x["impostos"] for x in p["linhas"] if x not in tri)
    # reforma tributária: 2026 = teste (CBS/IBS só informativos); 2027 = CBS no lugar de PIS/COFINS
    for x in p["linhas"]:
        ano = int(x["mes"][:4])
        if ano == 2026:
            assert "PIS" in x["impostos"] and "CBS" not in x["impostos"] and "CBS" in x["informativos"]
        elif ano in (2027, 2028):
            assert "PIS" not in x["impostos"] and "CBS" in x["impostos"] and "ISS" in x["impostos"]
    from app.modules.financas import impostos as imp

    assert imp.reforma(2030)["iss"] == Decimal("0.8") and imp.reforma(2033)["iss"] == 0
    assert cliente.delete(f"{F}/impostos/{eid}/receitas-informadas/{mes_passado}", headers=h).status_code == 200


def test_dre_mostra_transferencias_e_lucro_real(cliente):  # noqa: F811
    h = _login(cliente)
    a = _conta(cliente, h, "Conta A DRE", saldo="0")
    b = _conta(cliente, h, "Conta B DRE", saldo="0")
    hoje = date.today().isoformat()
    cliente.post(f"{F}/transferencias", headers=h, json={"origem_id": a["id"], "destino_id": b["id"], "data": hoje, "valor": "300"})
    tudo = cliente.get(f"{F}/dre", headers=h, params={"de": hoje, "ate": hoje}).json()
    tr = next(lin for lin in tudo["linhas"] if lin["chave"] == "transferencias")
    nomes = {i["categoria"]: i["total"] for i in tr["itens"]}
    assert nomes["Enviadas de Conta A DRE"] == -300 and nomes["Recebidas em Conta B DRE"] == 300 and tr["total"] == 0
    so_b = cliente.get(f"{F}/dre", headers=h, params={"de": hoje, "ate": hoje, "conta_ids": [b["id"]]}).json()
    assert next(lin for lin in so_b["linhas"] if lin["chave"] == "geracao_caixa")["total"] == 300

    p = cliente.get(f"{F}/impostos", headers=h).json()
    cfg = {**p["config"], "regime": "real"}
    cfg.pop("aliquota_iss_emissora")
    cfg.pop("regime_cadastro")
    p = cliente.put(f"{F}/impostos/{p['emissora_id']}/config", headers=h, json=cfg).json()
    assert p["resumo"]["regime"] == "real" and "Lucro Real" in p["aviso"]
    assert all("IRPJ" not in x["impostos"] for x in p["linhas"])
    emissora = next(e for e in cliente.get("/api/nfse/emissoras", headers=h).json() if e["id"] == p["emissora_id"])
    assert emissora["regime"] == "real"  # o regime mora no cadastro da empresa
    cfg["regime"] = "presumido"
    cliente.put(f"{F}/impostos/{p['emissora_id']}/config", headers=h, json=cfg)
