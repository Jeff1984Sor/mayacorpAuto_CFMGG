"""Planilhas do financeiro: modelos para importar (contas a pagar, extrato) e exportações (extrato, fluxo, DRE, contas)."""

import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "1E2D6B"
CABECALHO = PatternFill("solid", fgColor=NAVY)
FONTE_CAB = Font(bold=True, color="FFFFFF")
BRL = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
DATA = "DD/MM/YYYY"


def _wb_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def tabela(titulo: str, cabecalho: list[str], linhas: list[list], *, larguras: list[int] | None = None,
           moeda: set[int] | None = None, datas: set[int] | None = None, topo: list[str] | None = None,
           rodape: list | None = None) -> bytes:
    """Planilha simples: título/linhas de topo, cabeçalho navy, dados, rodapé de totais e autofiltro."""
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]
    r = 1
    for t in topo or []:
        ws.cell(r, 1, t).font = Font(bold=r == 1, size=13 if r == 1 else 10, color=NAVY if r == 1 else "444444")
        r += 1
    if topo:
        r += 1
    cab = r
    for i, h in enumerate(cabecalho, start=1):
        c = ws.cell(r, i, h)
        c.fill, c.font, c.alignment = CABECALHO, FONTE_CAB, Alignment(vertical="center")
    for linha in linhas:
        r += 1
        for i, v in enumerate(linha, start=1):
            c = ws.cell(r, i, float(v) if isinstance(v, Decimal) else v)
            if moeda and i - 1 in moeda:
                c.number_format = BRL
            if datas and i - 1 in datas:
                c.number_format = DATA
    if rodape:
        r += 1
        for i, v in enumerate(rodape, start=1):
            c = ws.cell(r, i, float(v) if isinstance(v, Decimal) else v)
            c.font = Font(bold=True)
            if moeda and i - 1 in moeda:
                c.number_format = BRL
    for i, w in enumerate(larguras or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A{cab}:{get_column_letter(len(cabecalho))}{max(cab, r - (1 if rodape else 0))}"
    ws.freeze_panes = ws.cell(cab + 1, 1)
    return _wb_bytes(wb)


# ---------------------------------------------------------------- modelo de contas a pagar
COLUNAS_CONTAS = ["Vencimento (dd/mm/aaaa)", "Pagamento (dd/mm/aaaa)", "Categoria", "Subcategoria", "Descrição", "Favorecido",
                  "Valor (R$)", "Forma de pagamento", "Documento", "Observação"]
FORMAS = {"pix": "pix", "boleto": "boleto", "ted": "ted", "ted/transferência": "ted", "transferência": "ted", "transferencia": "ted",
          "dinheiro": "dinheiro", "cartão": "cartao", "cartao": "cartao", "débito automático": "debito", "debito automatico": "debito",
          "débito": "debito", "debito": "debito", "outro": "outro"}


def modelo_contas_pagar() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Contas"
    for i, h in enumerate(COLUNAS_CONTAS, start=1):
        c = ws.cell(1, i, h)
        c.fill, c.font = PatternFill("solid", fgColor="F5A623"), Font(bold=True, color="0F1729")
        ws.column_dimensions[get_column_letter(i)].width = [16, 16, 24, 20, 32, 28, 14, 20, 16, 30][i - 1]
    hoje = date.today()
    ws.append([hoje.replace(day=20), hoje.replace(day=20), "IRPJ e CSLL", "IRPJ", "IRPJ trimestral", "Receita Federal", 1500.00,
               "Boleto", "DARF 2089", "Exemplo: já pago"])
    ws.append([hoje.replace(day=5), None, "Folha e encargos", "Salários", "Folha do mês", "Folha de pagamento", 25000.00, "TED", "", ""])
    for row in ws.iter_rows(min_row=2, max_row=3):
        row[0].number_format = row[1].number_format = DATA
        row[6].number_format = BRL
    inst = wb.create_sheet("Instruções")
    for linha in [
        "Como preencher (uma conta por linha, a partir da linha 2):",
        "• Vencimento: obrigatório. Aceita data do Excel, dd/mm/aaaa ou aaaa-mm-dd.",
        "• Pagamento: deixe vazio para conta em aberto. Preenchido = importa como paga (sem lançar no extrato).",
        "• Categoria: obrigatória. Se não existir, é criada como despesa. Subcategoria: opcional.",
        "• Descrição: se vazia, usa o nome da categoria. Favorecido: cria o fornecedor se ainda não existir.",
        "• Valor: obrigatório. Aceita 1500,00 / 1.500,00 / 1500.00.",
        "• Forma de pagamento: PIX, Boleto, TED/Transferência, Dinheiro, Cartão, Débito automático ou Outro.",
        "• Apague as duas linhas de exemplo antes de importar.",
    ]:
        inst.append([linha])
    inst.column_dimensions["A"].width = 110
    return _wb_bytes(wb)


def modelo_extrato(nome_conta: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Extrato"
    for i, h in enumerate(["Data", "Descrição", "Entrada (R$)", "Saída (R$)", "Categoria / Subcategoria"], start=1):
        c = ws.cell(1, i, h)
        c.fill, c.font = CABECALHO, FONTE_CAB
        ws.column_dimensions[get_column_letter(i)].width = [14, 40, 16, 16, 34][i - 1]
    ws.append([date.today(), "Exemplo de recebimento", 1000.00, None, "Honorários advocatícios"])
    ws.append([date.today(), "Exemplo de tarifa", None, 25.90, "Tarifas bancárias : Manutenção de conta"])
    inst = wb.create_sheet("Instruções")
    for linha in [f"Lançamentos para a conta {nome_conta}.", "• Uma linha por lançamento. Data e valor são obrigatórios.",
                  "• Use Entrada OU Saída (ou uma coluna Valor com negativo para saída).",
                  "• Categoria / Subcategoria no formato 'Categoria : Subcategoria' (opcional; criada se não existir).",
                  "• Apague as linhas de exemplo antes de importar."]:
        inst.append([linha])
    inst.column_dimensions["A"].width = 100
    return _wb_bytes(wb)


# ---------------------------------------------------------------- leitura
def ler_data(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def ler_valor(v) -> Decimal | None:
    """Número do Excel ou texto em BRL ("R$ 1.500,00", "(25,90)", "-25.9")."""
    if v is None or v == "":
        return None
    if isinstance(v, int | float | Decimal):
        return Decimal(str(v)).quantize(Decimal("0.01"))
    s = str(v).strip().replace("R$", "").replace(" ", "")
    negativo = s.startswith("(") and s.endswith(")") or s.startswith("-")
    s = s.strip("()-")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        d = Decimal(s).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None
    return -d if negativo else d


def linhas_planilha(conteudo: bytes, aba: str | None = None) -> list[tuple]:
    wb = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    ws = wb[aba] if aba and aba in wb.sheetnames else wb.active
    return list(ws.iter_rows(values_only=True))


def _norm(s) -> str:
    import unicodedata

    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z]", "", s)


def colunas_extrato(cabecalho: tuple) -> dict[str, int]:
    """Acha as colunas do extrato pelo nome do cabeçalho (ou usa a posição padrão)."""
    pos: dict[str, int] = {}
    for i, h in enumerate(cabecalho):
        n = _norm(h)
        if not n:
            continue
        if n.startswith("data") and "data" not in pos:
            pos["data"] = i
        elif any(x in n for x in ("descric", "histor", "lancament")) and "descricao" not in pos:
            pos["descricao"] = i
        elif any(x in n for x in ("entrada", "credito")):
            pos["entrada"] = i
        elif any(x in n for x in ("saida", "debito")):
            pos["saida"] = i
        elif n.startswith("valor"):
            pos["valor"] = i
        elif n == "tipo":
            pos["tipo"] = i
        elif "categ" in n:
            pos["categoria"] = i
        elif "documento" in n and "descricao" not in pos:
            pos["descricao"] = i
    if "data" not in pos:
        return {"data": 0, "descricao": 1, "entrada": 2, "saida": 3, "categoria": 4, "_sem_cabecalho": 1}
    return pos
