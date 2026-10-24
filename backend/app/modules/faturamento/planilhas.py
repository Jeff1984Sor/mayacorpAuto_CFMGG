"""Exportar / importar Contratos e Casos em Excel (mesmas colunas nos dois sentidos).

Na importação, a linha com "Nº" de um registro existente atualiza; sem número (ou número novo), cria.
Cliente, emissora, tabela de preço, área e responsável são achados pelo nome (ou CPF/CNPJ / e-mail).
"""

import io
from datetime import date
from decimal import Decimal

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_, select

from app.core.deps import Ctx, Operador
from app.core.documentos import so_digitos
from app.core.models import Usuario
from app.modules.faturamento.judiciario import INSTANCIAS
from app.modules.faturamento.models import (
    ROTULO_TIPO_CASO,
    AreaAtuacao,
    CampoAdicional,
    Caso,
    Contrato,
    StatusCaso,
    TabelaPreco,
    TipoCampo,
)
from app.modules.financas.planilhas import ler_data, ler_valor
from app.modules.rpa_nfse.models import Emissora, Tomador

router = APIRouter(prefix="/faturamento", tags=["Faturamento · Planilhas"])
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NAVY = PatternFill("solid", fgColor="1E2D6B")

COL_CONTRATO = ["Nº", "Cliente", "CPF/CNPJ do cliente", "Nome do contrato", "Objeto", "Início", "Fim", "Empresa emissora",
                "Tabela de preço", "Ativo", "Observação"]
COL_CASO = ["Nº", "Nº do contrato", "Título", "Tipo do caso", "Nº do processo", "Parte contrária", "Área", "Polo", "Instância", "Tribunal", "Vara", "Comarca",
            "Valor da causa", "Distribuição", "Responsável (e-mail)", "Situação", "Abertura", "Encerramento", "Tabela de preço",
            "Descrição"]
POLOS = {"ativo": "ativo", "autor": "ativo", "polo ativo": "ativo", "passivo": "passivo", "réu": "passivo", "reu": "passivo",
         "polo passivo": "passivo", "terceiro": "terceiro", "consultivo": "consultivo"}
ROTULO_POLO = {"ativo": "Ativo", "passivo": "Passivo", "terceiro": "Terceiro", "consultivo": "Consultivo"}


def _planilha(titulo: str, cab: list[str], linhas: list[list], instrucoes: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = titulo
    ws.append(cab)
    for i in range(1, len(cab) + 1):
        c = ws.cell(1, i)
        c.fill, c.font = NAVY, Font(bold=True, color="FFFFFF")
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(cab[i - 1]) + 6))
    for linha in linhas:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in linha])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, date):
                c.number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{max(1, ws.max_row)}"
    inst = wb.create_sheet("Instruções")
    for t in instrucoes:
        inst.append([t])
    inst.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _baixar(conteudo: bytes, nome: str) -> Response:
    return Response(conteudo, media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{nome}"'})


def _ler(conteudo: bytes, nome: str) -> tuple[list[str], list[list]]:
    if not nome.lower().endswith(".xlsx"):
        raise HTTPException(400, "Envie a planilha .xlsx (use o modelo ou a exportação).")
    try:
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, "Não consegui ler a planilha.") from e
    ws = wb.worksheets[0]
    linhas = [list(r) for r in ws.iter_rows(values_only=True)]
    if not linhas:
        raise HTTPException(400, "Planilha vazia.")
    return [str(c or "").strip().lower() for c in linhas[0]], linhas[1:]


def _col(cab: list[str], row: list, *nomes: str):
    for n in nomes:
        n = n.lower()
        if n in cab:
            i = cab.index(n)
            if i < len(row) and row[i] not in (None, ""):
                v = row[i]
                return v.strip() if isinstance(v, str) else v
    return None


def _sim(v) -> bool | None:
    if v is None:
        return None
    return str(v).strip().lower() in ("sim", "s", "true", "1", "ativo", "x", "yes")


def _por_nome(ctx, modelo, nome) -> object | None:
    if not nome:
        return None
    return ctx.db.scalar(select(modelo).where(modelo.tenant_id == ctx.tenant_id, func.lower(modelo.nome) == str(nome).strip().lower()))


def _sem_acento(s: str) -> str:
    import unicodedata

    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().strip().lower()


def _num(v) -> int | None:
    try:
        return int(str(v).strip().lstrip("#")) if v not in (None, "") else None
    except ValueError:
        return None


# ================================================================ contratos
def _linhas_contratos(ctx) -> list[list]:
    saida = []
    for c in ctx.db.scalars(select(Contrato).where(Contrato.tenant_id == ctx.tenant_id).order_by(Contrato.numero)):
        em = ctx.db.get(Emissora, c.emissora_id) if c.emissora_id else None
        tab = ctx.db.get(TabelaPreco, c.tabela_preco_id) if c.tabela_preco_id else None
        saida.append([c.numero, c.tomador.nome if c.tomador else "", c.tomador.cpf_cnpj if c.tomador else None, c.nome, c.objeto,
                      c.data_inicio, c.data_fim, em.razao_social if em else None, tab.nome if tab else None,
                      "Sim" if c.ativo else "Não", c.observacao])
    return saida


INSTR_CONTRATOS = [
    "Uma linha por contrato. A planilha exportada pode ser editada e importada de volta.",
    "• Nº: deixe vazio para criar um contrato novo; com o nº de um contrato existente, a linha atualiza esse contrato.",
    "• Cliente: obrigatório. O sistema procura pelo CPF/CNPJ (se informado) e depois pelo nome exato do cadastro de clientes.",
    "• Nome do contrato: obrigatório. Início/Fim: dd/mm/aaaa.",
    "• Empresa emissora e Tabela de preço: nome exatamente como cadastrado (opcional).",
    "• Ativo: Sim ou Não (vazio = Sim).",
]


@router.get("/contratos/exportar.xlsx")
def exportar_contratos(ctx: Ctx):
    return _baixar(_planilha("Contratos", COL_CONTRATO, _linhas_contratos(ctx), INSTR_CONTRATOS), f"contratos_{date.today():%Y%m%d}.xlsx")


@router.get("/contratos/modelo.xlsx")
def modelo_contratos(ctx: Ctx):
    exemplo = [[None, "Condomínio Edifício Exemplo", "11444777000161", "Consultivo mensal", "Assessoria condominial", date.today(),
                None, None, None, "Sim", "Linha de exemplo: apague antes de importar"]]
    return _baixar(_planilha("Contratos", COL_CONTRATO, exemplo, INSTR_CONTRATOS), "modelo_contratos.xlsx")


@router.post("/contratos/importar")
async def importar_contratos(ctx: Operador, arquivo: UploadFile = File(...)):
    cab, linhas = _ler(await arquivo.read(), arquivo.filename or "")
    criados = atualizados = 0
    erros: list[str] = []
    proximo = (ctx.db.scalar(select(func.max(Contrato.numero)).where(Contrato.tenant_id == ctx.tenant_id)) or 0) + 1
    for i, row in enumerate(linhas, start=2):
        if not any(v not in (None, "") for v in row):
            continue
        nome = _col(cab, row, "nome do contrato", "contrato", "nome")
        cli_nome, doc = _col(cab, row, "cliente"), so_digitos(str(_col(cab, row, "cpf/cnpj do cliente", "cpf/cnpj") or ""))
        tom = None
        if doc:
            tom = ctx.db.scalar(select(Tomador).where(Tomador.tenant_id == ctx.tenant_id, Tomador.cpf_cnpj == doc))
        if not tom and cli_nome:
            tom = _por_nome(ctx, Tomador, cli_nome)
        if not nome:
            erros.append(f"Linha {i}: falta o nome do contrato.")
            continue
        if not tom:
            erros.append(f"Linha {i}: cliente “{cli_nome or doc}” não encontrado no cadastro.")
            continue
        n = _num(_col(cab, row, "nº", "numero", "número"))
        c = ctx.db.scalar(select(Contrato).where(Contrato.tenant_id == ctx.tenant_id, Contrato.numero == n)) if n else None
        if c:
            atualizados += 1
        else:
            c = Contrato(tenant_id=ctx.tenant_id, numero=n if n and n >= proximo else proximo, tomador_id=tom.id, nome=nome)
            proximo = max(proximo, c.numero) + 1
            ctx.db.add(c)
            criados += 1
        c.tomador_id, c.nome = tom.id, str(nome)[:200]
        if (v := _col(cab, row, "objeto")) is not None:
            c.objeto = str(v)
        c.data_inicio = ler_data(_col(cab, row, "início", "inicio")) or c.data_inicio
        c.data_fim = ler_data(_col(cab, row, "fim")) or c.data_fim
        if em_nome := _col(cab, row, "empresa emissora", "emissora"):
            em = ctx.db.scalar(select(Emissora).where(Emissora.tenant_id == ctx.tenant_id, or_(
                func.lower(Emissora.razao_social) == str(em_nome).lower(), func.lower(Emissora.nome_fantasia) == str(em_nome).lower(),
                Emissora.cnpj == so_digitos(str(em_nome)))))
            if em:
                c.emissora_id = em.id
            else:
                erros.append(f"Linha {i}: emissora “{em_nome}” não encontrada (contrato salvo sem ela).")
        if tab_nome := _col(cab, row, "tabela de preço", "tabela"):
            tab = _por_nome(ctx, TabelaPreco, tab_nome)
            if tab:
                c.tabela_preco_id = tab.id
            else:
                erros.append(f"Linha {i}: tabela “{tab_nome}” não encontrada.")
        ativo = _sim(_col(cab, row, "ativo"))
        c.ativo = True if ativo is None else ativo
        if (v := _col(cab, row, "observação", "observacao")) is not None:
            c.observacao = str(v)
        ctx.db.flush()
    ctx.auditar("faturamento.contrato.importar", "fat_contrato", descricao=f"{criados} novos, {atualizados} atualizados, {len(erros)} avisos")
    ctx.db.commit()
    return {"criados": criados, "atualizados": atualizados, "erros": erros[:50], "qtd_erros": len(erros)}


# ================================================================ casos
def _campos_caso(ctx) -> list[CampoAdicional]:
    return list(ctx.db.scalars(select(CampoAdicional).where(CampoAdicional.tenant_id == ctx.tenant_id, CampoAdicional.entidade == "caso",
                                                            CampoAdicional.ativo).order_by(CampoAdicional.ordem, CampoAdicional.id)))


def _instr_casos(campos: list[CampoAdicional]) -> list[str]:
    return [
        "Uma linha por caso. A planilha exportada pode ser editada e importada de volta.",
        "• Nº: vazio cria um caso novo; com o nº de um caso existente, a linha atualiza esse caso.",
        "• Nº do contrato: obrigatório (o número que aparece em Cadastros › Contratos). Título: obrigatório.",
        "• Tipo do caso: " + ", ".join(ROTULO_TIPO_CASO.values()) + ".",
        "• Área: nome da área de atuação (criada se ainda não existir). Polo: Ativo, Passivo, Terceiro ou Consultivo.",
        "• Valor da causa: 150000,00. Datas: dd/mm/aaaa. Situação: Aberto ou Encerrado.",
        "• Responsável: e-mail (ou nome) de um profissional cadastrado.",
        "• As colunas depois de Descrição são os campos adicionais configurados." + (
            " Listas aceitam: " + "; ".join(f"{c.nome} = {', '.join(c.opcoes)}" for c in campos if c.tipo == TipoCampo.lista)
            if any(c.tipo == TipoCampo.lista for c in campos) else ""),
    ]


@router.get("/casos/exportar.xlsx")
def exportar_casos(ctx: Ctx):
    campos = _campos_caso(ctx)
    linhas = []
    for c in ctx.db.scalars(select(Caso).where(Caso.tenant_id == ctx.tenant_id).order_by(Caso.numero)):
        resp = ctx.db.get(Usuario, c.responsavel_id) if c.responsavel_id else None
        tab = ctx.db.get(TabelaPreco, c.tabela_preco_id) if c.tabela_preco_id else None
        extra = []
        for cp in campos:
            v = (c.campos or {}).get(str(cp.id))
            if cp.tipo == TipoCampo.sim_nao and v is not None:
                v = "Sim" if v else "Não"
            elif cp.tipo == TipoCampo.data and v:
                v = date.fromisoformat(v)
            elif cp.tipo in (TipoCampo.numero, TipoCampo.moeda) and v is not None:
                v = float(v)
            extra.append(v)
        linhas.append([c.numero, c.contrato.numero, c.titulo, ROTULO_TIPO_CASO.get(c.tipo) if c.tipo else None, c.numero_processo, c.parte_contraria, c.area.nome if c.area else None,
                       ROTULO_POLO.get(c.polo or ""), INSTANCIAS.get(c.instancia or ""), c.tribunal, c.vara, c.comarca, c.valor_causa, c.data_distribuicao,
                       resp.email if resp else None, "Aberto" if c.status == StatusCaso.aberto else "Encerrado", c.data_abertura,
                       c.data_encerramento, tab.nome if tab else None, c.descricao, *extra])
    return _baixar(_planilha("Casos", COL_CASO + [cp.nome for cp in campos], linhas, _instr_casos(campos)),
                   f"casos_{date.today():%Y%m%d}.xlsx")


@router.get("/casos/modelo.xlsx")
def modelo_casos(ctx: Ctx):
    campos = _campos_caso(ctx)
    exemplo = [[None, 1, "Ação de cobrança — exemplo", "Judicial Estadual", "0000000-00.2026.8.26.0100", "Fulano de Tal", "Cível", "Ativo", "1ª instância", "TJSP",
                "1ª Vara Cível", "São Paulo", 50000, date.today(), None, "Aberto", date.today(), None, None,
                "Linha de exemplo: apague antes de importar", *[None for _ in campos]]]
    return _baixar(_planilha("Casos", COL_CASO + [cp.nome for cp in campos], exemplo, _instr_casos(campos)), "modelo_casos.xlsx")


@router.post("/casos/importar")
async def importar_casos(ctx: Operador, arquivo: UploadFile = File(...)):
    cab, linhas = _ler(await arquivo.read(), arquivo.filename or "")
    campos = _campos_caso(ctx)
    criados = atualizados = 0
    erros: list[str] = []
    proximo = (ctx.db.scalar(select(func.max(Caso.numero)).where(Caso.tenant_id == ctx.tenant_id)) or 0) + 1
    for i, row in enumerate(linhas, start=2):
        if not any(v not in (None, "") for v in row):
            continue
        titulo = _col(cab, row, "título", "titulo")
        n_contrato = _num(_col(cab, row, "nº do contrato", "contrato"))
        contrato = ctx.db.scalar(select(Contrato).where(Contrato.tenant_id == ctx.tenant_id, Contrato.numero == n_contrato)) if n_contrato else None
        if not titulo:
            erros.append(f"Linha {i}: falta o título.")
            continue
        if not contrato:
            erros.append(f"Linha {i}: contrato nº {n_contrato or '—'} não encontrado.")
            continue
        n = _num(_col(cab, row, "nº", "numero", "número"))
        c = ctx.db.scalar(select(Caso).where(Caso.tenant_id == ctx.tenant_id, Caso.numero == n)) if n else None
        if c:
            atualizados += 1
        else:
            c = Caso(tenant_id=ctx.tenant_id, numero=n if n and n >= proximo else proximo, contrato_id=contrato.id, titulo=titulo,
                     data_abertura=date.today())
            proximo = max(proximo, c.numero) + 1
            ctx.db.add(c)
            criados += 1
        c.contrato_id, c.titulo = contrato.id, str(titulo)[:250]
        for attr, *nomes in [("numero_processo", "nº do processo", "processo"), ("parte_contraria", "parte contrária", "parte contraria"),
                             ("tribunal", "tribunal"), ("vara", "vara"), ("comarca", "comarca"), ("descricao", "descrição", "descricao")]:
            if (v := _col(cab, row, *nomes)) is not None:
                setattr(c, attr, str(v))
        if area_nome := _col(cab, row, "área", "area"):
            area = _por_nome(ctx, AreaAtuacao, area_nome)
            if not area:
                area = AreaAtuacao(tenant_id=ctx.tenant_id, nome=str(area_nome)[:100])
                ctx.db.add(area)
                ctx.db.flush()
            c.area_id = area.id
        if tipo := _col(cab, row, "tipo do caso", "tipo"):
            achado = next((t for t, rot in ROTULO_TIPO_CASO.items() if _sem_acento(rot) == _sem_acento(str(tipo))
                           or t.value == str(tipo).strip().lower()), None)
            if achado:
                c.tipo = achado
            else:
                erros.append(f"Linha {i}: tipo “{tipo}” inválido (use {', '.join(ROTULO_TIPO_CASO.values())}).")
        if inst := _col(cab, row, "instância", "instancia"):
            c.instancia = next((k for k, rot in INSTANCIAS.items() if _sem_acento(rot) == _sem_acento(str(inst)) or k == str(inst).lower()), c.instancia)
        if polo := _col(cab, row, "polo"):
            c.polo = POLOS.get(str(polo).strip().lower())
        if (v := ler_valor(_col(cab, row, "valor da causa"))) is not None:
            c.valor_causa = abs(v)
        for attr, nome in [("data_distribuicao", "distribuição"), ("data_abertura", "abertura"), ("data_encerramento", "encerramento")]:
            if d := ler_data(_col(cab, row, nome, nome.replace("ç", "c").replace("ã", "a"))):
                setattr(c, attr, d)
        if resp := _col(cab, row, "responsável (e-mail)", "responsável", "responsavel"):
            u = ctx.db.scalar(select(Usuario).where(Usuario.tenant_id == ctx.tenant_id, or_(
                func.lower(Usuario.email) == str(resp).lower(), func.lower(Usuario.nome) == str(resp).lower())))
            if u:
                c.responsavel_id = u.id
            else:
                erros.append(f"Linha {i}: responsável “{resp}” não encontrado.")
        if sit := _col(cab, row, "situação", "situacao", "status"):
            c.status = StatusCaso.encerrado if str(sit).lower().startswith("encerr") else StatusCaso.aberto
        if tab_nome := _col(cab, row, "tabela de preço", "tabela"):
            tab = _por_nome(ctx, TabelaPreco, tab_nome)
            if tab:
                c.tabela_preco_id = tab.id
            else:
                erros.append(f"Linha {i}: tabela “{tab_nome}” não encontrada.")
        valores = dict(c.campos or {})
        for cp in campos:
            v = _col(cab, row, cp.nome)
            if v is None:
                continue
            if cp.tipo == TipoCampo.sim_nao:
                valores[str(cp.id)] = bool(_sim(v))
            elif cp.tipo == TipoCampo.data:
                d = ler_data(v)
                if d:
                    valores[str(cp.id)] = d.isoformat()
            elif cp.tipo in (TipoCampo.numero, TipoCampo.moeda):
                num = ler_valor(v)
                if num is not None:
                    valores[str(cp.id)] = str(num)
            elif cp.tipo == TipoCampo.lista and str(v) not in cp.opcoes:
                erros.append(f"Linha {i}: “{v}” não é uma opção de {cp.nome}.")
            else:
                valores[str(cp.id)] = str(v)
        c.campos = valores
        ctx.db.flush()
    ctx.auditar("faturamento.caso.importar", "fat_caso", descricao=f"{criados} novos, {atualizados} atualizados, {len(erros)} avisos")
    ctx.db.commit()
    return {"criados": criados, "atualizados": atualizados, "erros": erros[:50], "qtd_erros": len(erros)}
