import asyncio
import io
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

from fastapi import APIRouter, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse, Response
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sse_starlette.sse import EventSourceResponse

from app.core import audit, certificados
from app.core.config import settings
from app.core.db import SessionLocal, agora
from app.core.deps import Admin, Contexto, Ctx, Operador
from app.core.documentos import cnpj_valido, documento_valido, so_digitos
from app.core.models import Certificado
from app.core.schemas import Mensagem, Pagina
from app.modules.rpa_nfse import fatura, raio_x, service, tasks
from app.modules.rpa_nfse.models import (
    Emissora,
    Lote,
    Nfse,
    ParametroFiscal,
    Rps,
    StatusLote,
    StatusNfse,
    StatusRps,
    Tomador,
)
from app.modules.rpa_nfse.prefeitura_sp import ClientePrefeitura, FalhaComunicacao, link_nota
from app.modules.rpa_nfse.schemas import (
    Alerta,
    AnexoEmail,
    CancelarIn,
    ConfigPrefeituraIn,
    ConfigPrefeituraOut,
    EmailPreviaOut,
    EmissoraIn,
    EmissoraOut,
    EnviarEmailIn,
    EnviarEmailsLoteIn,
    EnvioLoteOut,
    Indicador,
    LoteDetalhe,
    LoteOut,
    ModeloEmailOut,
    NfseDetalhe,
    NfseOut,
    PainelOut,
    ParametroFiscalIn,
    ParametroFiscalOut,
    ProgressoOut,
    RaioXOut,
    RpsOut,
    RpsPatch,
    TomadorIn,
    TomadorOut,
    ZipIn,
)

router = APIRouter(prefix="/nfse", tags=["NFS-e"])
config_router = APIRouter(prefix="/config", tags=["configurações"])

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _erro(e: service.RegraNegocio):
    return HTTPException(400, str(e))


def _do_tenant(ctx: Contexto, modelo, oid: int):
    obj = ctx.db.get(modelo, oid)
    if not obj or getattr(obj, "tenant_id", None) != ctx.tenant_id:
        raise HTTPException(404, "Registro não encontrado.")
    return obj


def _baixar(conteudo: bytes, nome: str, mime: str) -> Response:
    return Response(conteudo, media_type=mime, headers={"Content-Disposition": f'attachment; filename="{nome}"'})


# ======================================================================== emissoras
def _emissora_out(ctx, e: Emissora) -> EmissoraOut:
    ok, msg = service.saude_emissora(ctx.db, e)
    return EmissoraOut.model_validate(e).model_copy(update={"saude_ok": ok, "saude_msg": msg})


@router.get("/emissoras", response_model=list[EmissoraOut])
def listar_emissoras(ctx: Ctx):
    q = select(Emissora).where(Emissora.tenant_id == ctx.tenant_id).order_by(Emissora.ativa.desc(), Emissora.razao_social)
    return [_emissora_out(ctx, e) for e in ctx.db.scalars(q)]


def _validar_emissora(ctx, dados: EmissoraIn) -> None:
    if not cnpj_valido(dados.cnpj):
        raise HTTPException(400, "CNPJ da emissora inválido.")
    if not dados.ccm or len(dados.ccm) > 8:
        raise HTTPException(400, "Informe a Inscrição Municipal (CCM) com até 8 dígitos.")
    if dados.certificado_id:
        _do_tenant(ctx, Certificado, dados.certificado_id)


@router.post("/emissoras", response_model=EmissoraOut)
def criar_emissora(dados: EmissoraIn, ctx: Admin):
    _validar_emissora(ctx, dados)
    e = Emissora(tenant_id=ctx.tenant_id, **dados.model_dump())
    if not e.certificado_id:
        ativo = ctx.db.scalar(select(Certificado).where(Certificado.tenant_id == ctx.tenant_id, Certificado.ativo))
        e.certificado_id = ativo.id if ativo else None
    ctx.db.add(e)
    try:
        ctx.db.flush()
    except IntegrityError:
        raise HTTPException(409, "Já existe uma emissora com este CNPJ e CCM.") from None
    ctx.auditar("emissora.criar", "emissora", e.id, descricao=e.razao_social, depois=audit.snapshot(e))
    ctx.db.commit()
    return _emissora_out(ctx, e)


@router.get("/emissoras/{eid}", response_model=EmissoraOut)
def ler_emissora(eid: int, ctx: Ctx):
    return _emissora_out(ctx, _do_tenant(ctx, Emissora, eid))


@router.put("/emissoras/{eid}", response_model=EmissoraOut)
def salvar_emissora(eid: int, dados: EmissoraIn, ctx: Admin):
    e = _do_tenant(ctx, Emissora, eid)
    _validar_emissora(ctx, dados)
    antes = audit.snapshot(e)
    for k, v in dados.model_dump().items():
        setattr(e, k, v)
    ctx.db.flush()
    a, d = audit.diff(antes, audit.snapshot(e))
    ctx.auditar("emissora.alterar", "emissora", e.id, descricao=e.razao_social, antes=a, depois=d)
    ctx.db.commit()
    return _emissora_out(ctx, e)


@router.delete("/emissoras/{eid}", response_model=Mensagem)
def excluir_emissora(eid: int, ctx: Admin):
    e = _do_tenant(ctx, Emissora, eid)
    if ctx.db.scalar(select(Lote.id).where(Lote.emissora_id == e.id).limit(1)):
        e.ativa = False
        ctx.auditar("emissora.inativar", "emissora", e.id, descricao=e.razao_social)
        ctx.db.commit()
        return Mensagem(mensagem="A emissora tem histórico de notas, então foi apenas inativada.")
    ctx.auditar("emissora.excluir", "emissora", e.id, descricao=e.razao_social, antes=audit.snapshot(e))
    ctx.db.query(ParametroFiscal).filter(ParametroFiscal.emissora_id == e.id).delete()
    ctx.db.delete(e)
    ctx.db.commit()
    return Mensagem(mensagem="Emissora excluída.")


@router.post("/emissoras/ler-nota")
async def ler_nota_emissora(ctx: Admin, arquivo: UploadFile = File(...)):
    """Lê uma NFS-e já emitida pela empresa e devolve cadastro + configuração fiscal (nada é salvo aqui)."""
    from app.core import ia
    from app.modules.rpa_nfse import leitura_nota

    nome = (arquivo.filename or "").lower()
    conteudo = await arquivo.read()
    if len(conteudo) > 10 * 1024 * 1024:
        raise HTTPException(400, "Arquivo maior que 10 MB.")
    mime = arquivo.content_type or ""
    try:
        if nome.endswith(".xml") or conteudo.lstrip()[:5] == b"<?xml" or conteudo.lstrip()[:1] == b"<":
            resultado = leitura_nota.validar(leitura_nota.ler_xml(conteudo), "xml")
        else:
            provedor = ia.provedor_ativo(ctx.db, ctx.tenant_id)
            if nome.endswith(".pdf") or conteudo[:5] == b"%PDF-":
                texto = leitura_nota.texto_do_pdf(conteudo)
                if len(texto) >= 80:
                    bruto, fonte = leitura_nota.ler_com_ia(provedor, texto, None), "pdf"
                elif provedor.nome == "gemini":  # PDF escaneado: o Gemini lê o PDF como imagem
                    bruto, fonte = leitura_nota.ler_com_ia(provedor, None, ("application/pdf", conteudo)), "pdf"
                else:
                    raise leitura_nota.NotaIlegivel("Este PDF é uma imagem escaneada. Envie o XML da nota ou uma foto.")
            elif mime.startswith("image/") or nome.endswith((".png", ".jpg", ".jpeg", ".webp")):
                bruto, fonte = leitura_nota.ler_com_ia(provedor, None, (mime or "image/jpeg", conteudo)), "imagem"
            else:
                raise leitura_nota.NotaIlegivel("Envie a nota em XML (recomendado), PDF ou foto.")
            resultado = leitura_nota.validar(bruto, fonte)
    except leitura_nota.NotaIlegivel as e:
        raise HTTPException(400, str(e)) from e
    except ia.IAIndisponivel as e:
        raise HTTPException(400, f"{e} (O XML da nota é lido sem IA.)") from e
    ctx.auditar("emissora.ler_nota", "emissora", descricao=f"Nota lida ({resultado['fonte']}): {arquivo.filename}")
    ctx.db.commit()
    return resultado


class ParametrosPadraoIn(BaseModel):
    nbs: str | None = None
    class_trib: str | None = None
    cst_ibscbs: str | None = None
    ind_operacao: str | None = None
    finalidade: str | None = None
    ind_final: str | None = None


@router.post("/emissoras/{eid}/parametros-padrao", response_model=ParametroFiscalOut)
def aplicar_parametros_padrao(eid: int, dados: ParametrosPadraoIn, ctx: Admin):
    """Cria ou atualiza o parâmetro fiscal "padrão" da emissora (usado quando o arquivo da Lello não traz IBS/CBS)."""
    e = _do_tenant(ctx, Emissora, eid)
    p = ctx.db.scalar(select(ParametroFiscal).where(ParametroFiscal.emissora_id == e.id, ParametroFiscal.escopo == "padrao")
                      .order_by(ParametroFiscal.id))
    novo = p is None
    if novo:
        p = ParametroFiscal(tenant_id=ctx.tenant_id, emissora_id=e.id, escopo="padrao",
                            descricao="Padrão lido da nota fiscal (confira com o contador)")
        ctx.db.add(p)
    antes = audit.snapshot(p) if not novo else None
    for k, v in dados.model_dump(exclude_none=True).items():
        setattr(p, k, so_digitos(v) or None)
    p.ativo = True
    ctx.db.flush()
    ctx.auditar("parametro_fiscal.da_nota", "parametro_fiscal", p.id, descricao=e.razao_social, antes=antes,
                depois=audit.snapshot(p))
    _reavaliar_abertos(ctx)
    ctx.db.commit()
    return p


# ======================================================================== tomadores
def _tomador_out(ctx, t: Tomador, stats: tuple[int, Decimal] | None = None) -> TomadorOut:
    if stats is None:
        stats = ctx.db.execute(
            select(func.count(Nfse.id), func.coalesce(func.sum(Nfse.valor), 0)).where(
                Nfse.tomador_id == t.id, Nfse.status == StatusNfse.emitida
            )
        ).one()
    return TomadorOut.model_validate(t).model_copy(
        update={"incompleto": service.tomador_incompleto(t), "qtd_notas": stats[0], "valor_acumulado": Decimal(stats[1])}
    )


def _filtro_tomadores(ctx, q: str | None, incompleto: bool | None):
    f = [Tomador.tenant_id == ctx.tenant_id]
    if q:
        like = f"%{q}%"
        dig = so_digitos(q)
        conds = [Tomador.nome.ilike(like)]
        if dig:
            conds.append(Tomador.cpf_cnpj.like(f"%{dig}%"))
        f.append(or_(*conds))
    if incompleto:
        f.append(or_(Tomador.cpf_cnpj.is_(None), Tomador.cep.is_(None), Tomador.logradouro.is_(None),
                     func.coalesce(func.json_array_length(Tomador.emails), 0) == 0))
    return f


@router.get("/tomadores", response_model=Pagina[TomadorOut])
def listar_tomadores(
    ctx: Ctx, q: str | None = None, incompleto: bool | None = None,
    pagina: int = Query(1, ge=1), por_pagina: int = Query(25, ge=1, le=200),
):
    f = _filtro_tomadores(ctx, q, incompleto)
    total = ctx.db.scalar(select(func.count()).select_from(Tomador).where(*f)) or 0
    itens = list(ctx.db.scalars(select(Tomador).where(*f).order_by(Tomador.nome).offset((pagina - 1) * por_pagina).limit(por_pagina)))
    stats = {
        tid: (n, v)
        for tid, n, v in ctx.db.execute(
            select(Nfse.tomador_id, func.count(Nfse.id), func.coalesce(func.sum(Nfse.valor), 0))
            .where(Nfse.tomador_id.in_([t.id for t in itens]), Nfse.status == StatusNfse.emitida)
            .group_by(Nfse.tomador_id)
        )
    }
    return Pagina[TomadorOut](itens=[_tomador_out(ctx, t, stats.get(t.id, (0, 0))) for t in itens], total=total,
                              pagina=pagina, por_pagina=por_pagina)


def _validar_tomador(dados: TomadorIn) -> None:
    if dados.cpf_cnpj and not documento_valido(dados.cpf_cnpj):
        raise HTTPException(400, "CPF/CNPJ inválido.")
    if dados.cpf_cnpj and dados.tipo == "PF" and len(dados.cpf_cnpj) != 11:
        raise HTTPException(400, "Pessoa física precisa de CPF (11 dígitos).")
    if dados.cpf_cnpj and dados.tipo == "PJ" and len(dados.cpf_cnpj) != 14:
        raise HTTPException(400, "Pessoa jurídica precisa de CNPJ (14 dígitos).")


def _dados_tomador(dados: TomadorIn) -> dict:
    """Listas viram JSON; o telefone principal acompanha o 1º da lista; só um contato principal."""
    d = dados.model_dump(mode="json")
    if d["telefones"]:
        d["telefone"] = d["telefones"][0]["numero"]
    if d["contatos"]:
        if not any(c["principal"] for c in d["contatos"]):
            d["contatos"][0]["principal"] = True
        visto = False
        for c in d["contatos"]:
            c["principal"], visto = (c["principal"] and not visto), visto or c["principal"]
    return d


@router.post("/tomadores", response_model=TomadorOut)
def criar_tomador(dados: TomadorIn, ctx: Operador):
    _validar_tomador(dados)
    t = Tomador(tenant_id=ctx.tenant_id, **_dados_tomador(dados))
    ctx.db.add(t)
    try:
        ctx.db.flush()
    except IntegrityError:
        raise HTTPException(409, "Já existe um cliente com este CPF/CNPJ.") from None
    ctx.auditar("tomador.criar", "tomador", t.id, descricao=t.nome, depois=audit.snapshot(t))
    ctx.db.commit()
    return _tomador_out(ctx, t)


@router.get("/tomadores/exportar")
def exportar_tomadores(ctx: Ctx):
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    cab = ["Tipo", "Nome", "CPF/CNPJ", "IM", "IE", "Logradouro", "Número", "Complemento", "Bairro", "Cidade", "UF",
           "CEP", "E-mails", "Telefone", "Observações"]
    ws.append(cab)
    for t in ctx.db.scalars(select(Tomador).where(Tomador.tenant_id == ctx.tenant_id).order_by(Tomador.nome)):
        ws.append([t.tipo, t.nome, t.cpf_cnpj, t.im, t.ie, t.logradouro, t.numero, t.complemento, t.bairro, t.cidade,
                   t.uf, t.cep, "; ".join(t.emails or []), t.telefone, t.observacoes])
    buf = io.BytesIO()
    wb.save(buf)
    return _baixar(buf.getvalue(), "clientes.xlsx", XLSX)


@router.post("/tomadores/importar", response_model=Mensagem)
async def importar_tomadores(ctx: Operador, arquivo: UploadFile = File(...)):
    conteudo = await arquivo.read()
    nome = (arquivo.filename or "").lower()
    linhas: list[list] = []
    if nome.endswith(".xlsx"):
        ws = load_workbook(io.BytesIO(conteudo), read_only=True).active
        linhas = [list(r) for r in ws.iter_rows(values_only=True)]
    elif nome.endswith(".csv"):
        import csv

        texto = conteudo.decode("utf-8-sig", errors="replace")
        sep = ";" if texto.count(";") > texto.count(",") else ","
        linhas = list(csv.reader(io.StringIO(texto), delimiter=sep))
    else:
        raise HTTPException(400, "Envie uma planilha .xlsx ou .csv (mesmo modelo da exportação).")
    if not linhas:
        raise HTTPException(400, "Planilha vazia.")
    cab = [str(c or "").strip().lower() for c in linhas[0]]

    def col(row, *nomes):
        for n in nomes:
            if n in cab and cab.index(n) < len(row):
                v = row[cab.index(n)]
                return str(v).strip() if v not in (None, "") else None
        return None

    criados = atualizados = ignorados = 0
    for row in linhas[1:]:
        nome_t = col(row, "nome", "razão social", "razao social")
        doc = so_digitos(col(row, "cpf/cnpj", "cpf_cnpj", "documento"))
        if not nome_t or (doc and not documento_valido(doc)):
            ignorados += 1
            continue
        t = ctx.db.scalar(select(Tomador).where(Tomador.tenant_id == ctx.tenant_id, Tomador.cpf_cnpj == doc)) if doc else None
        if t is None:
            t = Tomador(tenant_id=ctx.tenant_id, tipo="PF" if len(doc) == 11 else "PJ", nome=nome_t, cpf_cnpj=doc or None, emails=[])
            ctx.db.add(t)
            criados += 1
        else:
            atualizados += 1
        for campo, *nomes in [("im", "im"), ("ie", "ie"), ("logradouro", "logradouro"), ("numero", "número", "numero"),
                              ("complemento", "complemento"), ("bairro", "bairro"), ("cidade", "cidade"), ("uf", "uf"),
                              ("telefone", "telefone"), ("observacoes", "observações", "observacoes")]:
            if (v := col(row, *nomes)) is not None:
                setattr(t, campo, v)
        if cep := so_digitos(col(row, "cep")):
            t.cep = cep
        if emails := col(row, "e-mails", "emails", "e-mail", "email"):
            t.emails = sorted({e.strip().lower() for e in emails.replace(",", ";").split(";") if "@" in e} | set(t.emails or []))
        t.nome = nome_t
    ctx.auditar("tomador.importar", "tomador", descricao=f"{criados} novos, {atualizados} atualizados, {ignorados} ignorados")
    ctx.db.commit()
    return Mensagem(mensagem=f"{criados} clientes novos, {atualizados} atualizados, {ignorados} linhas ignoradas.")


@router.get("/tomadores/{tid}", response_model=TomadorOut)
def ler_tomador(tid: int, ctx: Ctx):
    return _tomador_out(ctx, _do_tenant(ctx, Tomador, tid))


@router.put("/tomadores/{tid}", response_model=TomadorOut)
def salvar_tomador(tid: int, dados: TomadorIn, ctx: Operador):
    t = _do_tenant(ctx, Tomador, tid)
    _validar_tomador(dados)
    antes = audit.snapshot(t)
    for k, v in _dados_tomador(dados).items():
        setattr(t, k, v)
    try:
        ctx.db.flush()
    except IntegrityError:
        raise HTTPException(409, "Já existe um cliente com este CPF/CNPJ.") from None
    a, d = audit.diff(antes, audit.snapshot(t))
    ctx.auditar("tomador.alterar", "tomador", t.id, descricao=t.nome, antes=a, depois=d)
    ctx.db.commit()
    return _tomador_out(ctx, t)


@router.delete("/tomadores/{tid}", response_model=Mensagem)
def excluir_tomador(tid: int, ctx: Admin):
    t = _do_tenant(ctx, Tomador, tid)
    if ctx.db.scalar(select(Nfse.id).where(Nfse.tomador_id == t.id).limit(1)):
        raise HTTPException(409, "Este cliente tem notas emitidas e não pode ser excluído.")
    ctx.db.query(Rps).filter(Rps.tomador_id == t.id).update({Rps.tomador_id: None})
    ctx.auditar("tomador.excluir", "tomador", t.id, descricao=t.nome, antes=audit.snapshot(t))
    ctx.db.delete(t)
    ctx.db.commit()
    return Mensagem(mensagem="Cliente excluído.")


@router.get("/tomadores/{tid}/notas", response_model=list[NfseOut])
def notas_do_tomador(tid: int, ctx: Ctx):
    _do_tenant(ctx, Tomador, tid)
    q = select(Nfse).where(Nfse.tomador_id == tid).order_by(Nfse.data_emissao.desc()).limit(200)
    return [_nfse_out(n) for n in ctx.db.scalars(q)]


# ======================================================================== parâmetros fiscais
@config_router.get("/parametros-fiscais", response_model=list[ParametroFiscalOut])
def listar_parametros(ctx: Ctx, emissora_id: int | None = None):
    f = [ParametroFiscal.tenant_id == ctx.tenant_id]
    if emissora_id:
        f.append(ParametroFiscal.emissora_id == emissora_id)
    ordem = {"padrao": 0, "servico": 1, "tomador": 2}
    itens = list(ctx.db.scalars(select(ParametroFiscal).where(*f)))
    itens.sort(key=lambda p: (p.emissora_id, ordem.get(p.escopo, 9), p.prioridade, p.id))
    return itens


def _validar_parametro(ctx, dados: ParametroFiscalIn) -> None:
    _do_tenant(ctx, Emissora, dados.emissora_id)
    if dados.escopo == "servico" and not dados.cod_servico:
        raise HTTPException(400, "Regra por serviço precisa do código de serviço.")
    if dados.escopo == "tomador":
        if not dados.tomador_id:
            raise HTTPException(400, "Regra por cliente precisa do cliente.")
        _do_tenant(ctx, Tomador, dados.tomador_id)
    for campo, tam in (("nbs", 9), ("class_trib", 6), ("ind_operacao", 6)):
        v = so_digitos(getattr(dados, campo))
        if getattr(dados, campo) and len(v) != tam:
            raise HTTPException(400, f"{campo} deve ter {tam} dígitos.")


def _reavaliar_abertos(ctx) -> None:
    """Parâmetros mudaram: lotes ainda em conferência são revalidados."""
    for lote in ctx.db.scalars(select(Lote).where(Lote.tenant_id == ctx.tenant_id, Lote.status == StatusLote.validando)):
        service.revalidar(ctx.db, lote)


@config_router.post("/parametros-fiscais", response_model=ParametroFiscalOut)
def criar_parametro(dados: ParametroFiscalIn, ctx: Admin):
    _validar_parametro(ctx, dados)
    p = ParametroFiscal(tenant_id=ctx.tenant_id, **dados.model_dump())
    ctx.db.add(p)
    ctx.db.flush()
    ctx.auditar("parametro_fiscal.criar", "parametro_fiscal", p.id, descricao=p.descricao, depois=audit.snapshot(p))
    _reavaliar_abertos(ctx)
    ctx.db.commit()
    return p


@config_router.put("/parametros-fiscais/{pid}", response_model=ParametroFiscalOut)
def salvar_parametro(pid: int, dados: ParametroFiscalIn, ctx: Admin):
    p = _do_tenant(ctx, ParametroFiscal, pid)
    _validar_parametro(ctx, dados)
    antes = audit.snapshot(p)
    for k, v in dados.model_dump().items():
        setattr(p, k, v)
    ctx.db.flush()
    a, d = audit.diff(antes, audit.snapshot(p))
    ctx.auditar("parametro_fiscal.alterar", "parametro_fiscal", p.id, descricao=p.descricao, antes=a, depois=d)
    _reavaliar_abertos(ctx)
    ctx.db.commit()
    return p


@config_router.delete("/parametros-fiscais/{pid}", response_model=Mensagem)
def excluir_parametro(pid: int, ctx: Admin):
    p = _do_tenant(ctx, ParametroFiscal, pid)
    ctx.auditar("parametro_fiscal.excluir", "parametro_fiscal", p.id, descricao=p.descricao, antes=audit.snapshot(p))
    ctx.db.delete(p)
    ctx.db.flush()
    _reavaliar_abertos(ctx)
    ctx.db.commit()
    return Mensagem(mensagem="Parâmetro excluído.")


# ======================================================================== prefeitura
def _cfg_prefeitura_out(ctx) -> ConfigPrefeituraOut:
    modo = service.modo_prefeitura(ctx.db, ctx.tenant_id)
    return ConfigPrefeituraOut(
        modo=modo, modos_permitidos=service.modos_permitidos(), ambiente=settings.ambiente,
        endpoint="—" if modo == "simulado" else settings.prefeitura_url_v2,
    )


@config_router.get("/prefeitura", response_model=ConfigPrefeituraOut)
def ler_prefeitura(ctx: Ctx):
    return _cfg_prefeitura_out(ctx)


@config_router.put("/prefeitura", response_model=ConfigPrefeituraOut)
def salvar_prefeitura(dados: ConfigPrefeituraIn, ctx: Admin):
    if dados.modo not in service.modos_permitidos():
        raise HTTPException(400, "Emissão real só pode ser ligada no servidor de produção.")
    if dados.modo == "producao" and (dados.confirmacao or "").strip().upper() != "EMITIR NOTAS REAIS":
        raise HTTPException(400, "Para ligar a emissão real, digite EMITIR NOTAS REAIS na confirmação.")
    if dados.modo != "simulado" and not ctx.db.scalar(
        select(Certificado.id).where(Certificado.tenant_id == ctx.tenant_id, Certificado.ativo)
    ):
        raise HTTPException(400, "Envie e ative o certificado digital A1 antes de sair do modo simulação.")
    cfg = service.config_nfse(ctx.db, ctx.tenant_id)
    antes = cfg.prefeitura_modo
    cfg.prefeitura_modo = dados.modo
    ctx.auditar("prefeitura.modo", "config_nfse", ctx.tenant_id, descricao=f"Modo da Prefeitura: {antes} → {dados.modo}",
                antes={"modo": antes}, depois={"modo": dados.modo})
    ctx.db.commit()
    return _cfg_prefeitura_out(ctx)


# ======================================================================== importação
def _rps_out(r: Rps) -> RpsOut:
    tom = (r.dados_json or {}).get("tomador", {})
    return RpsOut(
        id=r.id, linha=r.linha, selecionado=r.selecionado, serie=r.serie, numero=r.numero, data_emissao=r.data_emissao,
        tomador_id=r.tomador_id, tomador_nome=tom.get("nome") or "", tomador_doc=tom.get("cpf_cnpj"),
        tomador_novo=r.tomador_novo, cod_servico=r.cod_servico, aliquota=r.aliquota, valor=r.valor, deducoes=r.deducoes,
        iss=r.iss, iss_retido=r.iss_retido, discriminacao=r.discriminacao, status=r.status, erros=r.erros_json or [],
        dados={"tomador": tom, "layout2": (r.dados_json or {}).get("layout2", {})},
        nfse_id=r.nfse.id if r.nfse else None, nfse_numero=r.nfse.numero if r.nfse else None,
    )


def _lote_out(lote: Lote) -> LoteOut:
    from app.core.models import Usuario

    usuario = None
    if lote.usuario_id:
        from sqlalchemy.orm import object_session

        u = object_session(lote).get(Usuario, lote.usuario_id)
        usuario = u.nome if u else None
    return LoteOut(
        id=lote.id, numero=lote.numero, emissora_id=lote.emissora_id,
        emissora_nome=lote.emissora.nome_fantasia or lote.emissora.razao_social, usuario_nome=usuario,
        arquivo_origem=lote.arquivo_origem, qtd=lote.qtd, valor_total=lote.valor_total, iss_total=lote.iss_total,
        status=lote.status, sublotes_total=lote.sublotes_total, sublotes_feitos=lote.sublotes_feitos,
        qtd_emitidas=lote.qtd_emitidas, qtd_rejeitadas=lote.qtd_rejeitadas, mensagem=lote.mensagem,
        criado_em=lote.criado_em, iniciado_em=lote.iniciado_em, concluido_em=lote.concluido_em,
    )


def _lote_detalhe(lote: Lote) -> LoteDetalhe:
    return LoteDetalhe(
        **_lote_out(lote).model_dump(), encoding=lote.encoding, eventos=lote.eventos or [],
        resumo=service.resumo(lote), rps=[_rps_out(r) for r in lote.rps if r.status != StatusRps.descartado],
    )


@router.post("/importacoes", response_model=LoteDetalhe)
async def importar(ctx: Operador, arquivo: UploadFile = File(...), emissora_id: int = Form(...)):
    emissora = _do_tenant(ctx, Emissora, emissora_id)
    if not emissora.ativa:
        raise HTTPException(400, "Esta emissora está inativa.")
    nome = arquivo.filename or "arquivo.txt"
    if Path(nome).suffix.lower() not in (".txt", ".rps", ""):
        raise HTTPException(400, "Por enquanto aceitamos o arquivo .txt de RPS da Prefeitura (layout gerado pela Lello).")
    conteudo = await arquivo.read()
    if len(conteudo) > 20 * 1024 * 1024:
        raise HTTPException(400, "Arquivo maior que 20 MB.")
    try:
        lote = service.importar(ctx.db, tenant_id=ctx.tenant_id, usuario_id=ctx.usuario.id, emissora=emissora,
                                nome=nome, conteudo=conteudo)
    except service.RegraNegocio as e:
        raise _erro(e) from e
    ctx.auditar("lote.importar", "lote", lote.id, descricao=f"Lote #{lote.numero}: {nome} ({lote.qtd} RPS)")
    ctx.db.commit()
    return _lote_detalhe(lote)


@router.get("/importacoes/{lid}", response_model=LoteDetalhe)
def ler_importacao(lid: int, ctx: Ctx):
    return _lote_detalhe(_do_tenant(ctx, Lote, lid))


@router.get("/importacoes/{lid}/raio-x", response_model=RaioXOut)
def raio_x_importacao(lid: int, ctx: Ctx):
    return raio_x.gerar(ctx.db, _do_tenant(ctx, Lote, lid))


@router.patch("/importacoes/{lid}/rps/{rid}")
def corrigir(lid: int, rid: int, dados: RpsPatch, ctx: Operador):
    lote = _do_tenant(ctx, Lote, lid)
    rps = ctx.db.get(Rps, rid)
    if not rps or rps.lote_id != lote.id:
        raise HTTPException(404, "RPS não encontrado.")
    antes = {"valor": str(rps.valor), "tomador": (rps.dados_json or {}).get("tomador"), "status": rps.status.value}
    try:
        service.corrigir_rps(ctx.db, lote, rps, dados.model_dump(exclude_unset=True))
    except service.RegraNegocio as e:
        raise _erro(e) from e
    depois = {"valor": str(rps.valor), "tomador": (rps.dados_json or {}).get("tomador"), "status": rps.status.value}
    ctx.auditar("rps.corrigir", "rps", rps.id, descricao=f"Lote #{lote.numero} · RPS {rps.numero}", antes=antes, depois=depois)
    ctx.db.commit()
    return {"rps": _rps_out(rps), "resumo": service.resumo(lote)}


@router.post("/importacoes/{lid}/selecao")
def selecao(lid: int, body: dict, ctx: Operador):
    lote = _do_tenant(ctx, Lote, lid)
    if lote.status != StatusLote.validando:
        raise HTTPException(400, "Este lote já foi enviado.")
    service.selecionar(lote, bool(body.get("selecionado")), body.get("status"))
    ctx.db.commit()
    return service.resumo(lote)


@router.post("/importacoes/{lid}/emitir", response_model=LoteOut)
def emitir(lid: int, ctx: Operador):
    lote = _do_tenant(ctx, Lote, lid)
    if lote.status != StatusLote.validando:
        raise HTTPException(400, "Este lote já foi enviado.")
    try:
        n = service.enfileirar(ctx.db, lote)
        ctx.db.flush()
    except service.RegraNegocio as e:
        ctx.db.rollback()
        raise _erro(e) from e
    except IntegrityError:
        ctx.db.rollback()
        raise HTTPException(409, "Algum destes RPS já foi emitido ou está na fila em outro lote.") from None
    ctx.auditar("lote.emitir", "lote", lote.id, descricao=f"Lote #{lote.numero}: {n} RPS enviados para emissão")
    ctx.db.commit()
    tasks.enfileirar_job(lote.id)
    return _lote_out(lote)


@router.delete("/importacoes/{lid}", response_model=Mensagem)
def descartar(lid: int, ctx: Operador):
    lote = _do_tenant(ctx, Lote, lid)
    if lote.status != StatusLote.validando:
        raise HTTPException(400, "Só dá para descartar um lote que ainda não foi enviado.")
    lote.status = StatusLote.descartado
    lote.evento("descarte", f"Lote descartado por {ctx.usuario.nome}.")
    ctx.auditar("lote.descartar", "lote", lote.id, descricao=f"Lote #{lote.numero}")
    ctx.db.commit()
    return Mensagem(mensagem="Importação descartada.")


# ======================================================================== lotes
@router.get("/lotes", response_model=Pagina[LoteOut])
def listar_lotes(
    ctx: Ctx, q: str | None = None, status: StatusLote | None = None, emissora_id: int | None = None,
    de: date | None = None, ate: date | None = None,
    pagina: int = Query(1, ge=1), por_pagina: int = Query(25, ge=1, le=200),
):
    f = [Lote.tenant_id == ctx.tenant_id, Lote.status != StatusLote.descartado]
    if q:
        f.append(or_(Lote.arquivo_origem.ilike(f"%{q}%"), Lote.numero == (int(q) if q.isdigit() else -1)))
    if status:
        f.append(Lote.status == status)
    if emissora_id:
        f.append(Lote.emissora_id == emissora_id)
    if de:
        f.append(Lote.criado_em >= datetime.combine(de, time.min))
    if ate:
        f.append(Lote.criado_em <= datetime.combine(ate, time.max))
    total = ctx.db.scalar(select(func.count()).select_from(Lote).where(*f)) or 0
    itens = ctx.db.scalars(select(Lote).where(*f).order_by(Lote.id.desc()).offset((pagina - 1) * por_pagina).limit(por_pagina))
    return Pagina[LoteOut](itens=[_lote_out(lt) for lt in itens], total=total, pagina=pagina, por_pagina=por_pagina)


@router.get("/lotes/{lid}", response_model=LoteDetalhe)
def ler_lote(lid: int, ctx: Ctx):
    return _lote_detalhe(_do_tenant(ctx, Lote, lid))


def _progresso(lote: Lote) -> ProgressoOut:
    feitos = lote.qtd_emitidas + lote.qtd_rejeitadas
    alvo = max(sum(1 for r in lote.rps if r.status not in (StatusRps.descartado, StatusRps.ok, StatusRps.aviso, StatusRps.erro)), 1)
    pct = 100.0 if lote.status in (StatusLote.concluido, StatusLote.parcial, StatusLote.falhou) else round(min(feitos / alvo, 1) * 100, 1)
    return ProgressoOut(
        lote_id=lote.id, status=lote.status, sublotes_total=lote.sublotes_total, sublotes_feitos=lote.sublotes_feitos,
        qtd=alvo, qtd_emitidas=lote.qtd_emitidas, qtd_rejeitadas=lote.qtd_rejeitadas, percentual=pct,
        mensagem=lote.mensagem, eventos=(lote.eventos or [])[-8:],
    )


@router.get("/lotes/{lid}/progresso")
async def progresso(lid: int, request: Request, ctx: Ctx):
    _do_tenant(ctx, Lote, lid)

    async def eventos():
        ultimo = None
        while True:
            if await request.is_disconnected():
                return
            with SessionLocal() as db:
                lote = db.get(Lote, lid)
                p = _progresso(lote)
            dado = p.model_dump_json()
            if dado != ultimo:
                yield {"event": "progresso", "data": dado}
                ultimo = dado
            if p.status not in (StatusLote.na_fila, StatusLote.transmitindo):
                yield {"event": "fim", "data": dado}
                return
            await asyncio.sleep(1)

    return EventSourceResponse(eventos(), ping=15)


@router.get("/lotes/{lid}/arquivo")
def arquivo_original(lid: int, ctx: Ctx):
    lote = _do_tenant(ctx, Lote, lid)
    if not lote.arquivo_path or not Path(lote.arquivo_path).exists():
        raise HTTPException(404, "Arquivo original não encontrado.")
    return FileResponse(lote.arquivo_path, filename=lote.arquivo_origem, media_type="text/plain")


@router.get("/lotes/{lid}/conferencia")
def conferencia(lid: int, ctx: Ctx):
    lote = _do_tenant(ctx, Lote, lid)
    return _baixar(service.planilha_conferencia(lote), f"conferencia_lote_{lote.numero}.xlsx", XLSX)


@router.post("/lotes/{lid}/reprocessar-rejeitados", response_model=LoteOut)
def reprocessar(lid: int, ctx: Operador):
    lote = _do_tenant(ctx, Lote, lid)
    if lote.status not in (StatusLote.parcial, StatusLote.falhou):
        raise HTTPException(400, "Só lotes parciais ou com falha podem ser reprocessados.")
    try:
        # RPS presos na fila (Prefeitura fora do ar) também voltam
        presos = [r for r in lote.rps if r.status == StatusRps.na_fila]
        n = len(presos)
        if presos:
            lote.status = StatusLote.na_fila
        pendentes = (StatusRps.rejeitado, StatusRps.ok, StatusRps.aviso)
        if any(r.status in pendentes and not r.nfse for r in lote.rps):
            n += service.reabrir_rejeitados(ctx.db, lote)
        elif not presos:
            raise service.RegraNegocio("Não há RPS rejeitados para reprocessar.")
        ctx.db.flush()
    except service.RegraNegocio as e:
        ctx.db.rollback()
        raise _erro(e) from e
    ctx.auditar("lote.reprocessar", "lote", lote.id, descricao=f"Lote #{lote.numero}: {n} RPS reenviados")
    ctx.db.commit()
    tasks.enfileirar_job(lote.id)
    return _lote_out(lote)


@router.post("/lotes/{lid}/enviar-emails", response_model=Mensagem)
def enviar_emails_lote(lid: int, ctx: Operador):
    lote = _do_tenant(ctx, Lote, lid)
    ok = falhas = sem_email = 0
    for nota in service.notas_do_lote(ctx.db, lote.id):
        try:
            envio = service.enviar_email_nota(ctx.db, nota)
            ok += envio.status == "enviado"
            falhas += envio.status != "enviado"
        except service.RegraNegocio:
            sem_email += 1
    ctx.auditar("lote.enviar_emails", "lote", lote.id, descricao=f"{ok} enviados, {falhas} falhas, {sem_email} sem e-mail")
    ctx.db.commit()
    return Mensagem(mensagem=f"{ok} e-mails enviados" + (f", {falhas} falharam" if falhas else "")
                    + (f", {sem_email} clientes sem e-mail" if sem_email else "") + ".")


# ======================================================================== notas
def _nfse_out(n: Nfse) -> NfseOut:
    tom = (n.rps.dados_json or {}).get("tomador", {})
    return NfseOut(
        id=n.id, numero=n.numero, cod_verificacao=n.cod_verificacao, data_emissao=n.data_emissao,
        rps_serie=n.rps.serie, rps_numero=n.rps.numero, lote_id=n.rps.lote_id, emissora_id=n.emissora_id,
        tomador_id=n.tomador_id, tomador_nome=tom.get("nome") or (n.tomador.nome if n.tomador else ""),
        tomador_doc=tom.get("cpf_cnpj"), valor=n.valor, iss=n.iss, status=n.status, cancelada_em=n.cancelada_em,
        motivo_cancelamento=n.motivo_cancelamento, link_prefeitura=link_nota(n.emissora.ccm, n.numero, n.cod_verificacao),
        simulada=n.simulada, pago_em=n.pago_em, valor_pago=n.valor_pago, baixa_origem=n.baixa_origem,
        fatura_id=fat.id if (fat := _fatura_da_nota(n)) else None, fatura_codigo=fat.codigo if fat else None,
    )


def _fatura_da_nota(n: Nfse):
    from sqlalchemy.orm import object_session

    from app.modules.rpa_nfse.models import Fatura

    return object_session(n).scalar(select(Fatura).where(Fatura.nfse_id == n.id))


def _filtro_notas(ctx, q, status, emissora_id, lote_id, de, ate, valor_min, valor_max, pagamento=None):
    f = [Nfse.tenant_id == ctx.tenant_id]
    if pagamento == "paga":
        f.append(Nfse.pago_em.is_not(None))
    elif pagamento == "aberta":
        f += [Nfse.pago_em.is_(None), Nfse.status == StatusNfse.emitida]
    if q:
        dig = so_digitos(q)
        conds = [Tomador.nome.ilike(f"%{q}%")]
        if dig:
            conds += [Nfse.numero == int(dig), Rps.numero == int(dig), Tomador.cpf_cnpj.like(f"%{dig}%")]
        f.append(or_(*conds))
    if status:
        f.append(Nfse.status == status)
    if emissora_id:
        f.append(Nfse.emissora_id == emissora_id)
    if lote_id:
        f.append(Rps.lote_id == lote_id)
    if de:
        f.append(Nfse.data_emissao >= datetime.combine(de, time.min))
    if ate:
        f.append(Nfse.data_emissao <= datetime.combine(ate, time.max))
    if valor_min is not None:
        f.append(Nfse.valor >= valor_min)
    if valor_max is not None:
        f.append(Nfse.valor <= valor_max)
    return f


def _q_notas():
    return select(Nfse).join(Rps, Nfse.rps_id == Rps.id).outerjoin(Tomador, Nfse.tomador_id == Tomador.id)


@router.get("/notas", response_model=Pagina[NfseOut])
def listar_notas(
    ctx: Ctx, q: str | None = None, status: StatusNfse | None = None, emissora_id: int | None = None,
    lote_id: int | None = None, de: date | None = None, ate: date | None = None,
    valor_min: Decimal | None = None, valor_max: Decimal | None = None,
    pagamento: str | None = Query(None, pattern="^(paga|aberta)$"),
    pagina: int = Query(1, ge=1), por_pagina: int = Query(25, ge=1, le=200),
):
    f = _filtro_notas(ctx, q, status, emissora_id, lote_id, de, ate, valor_min, valor_max, pagamento)
    total = ctx.db.scalar(select(func.count()).select_from(_q_notas().where(*f).subquery())) or 0
    itens = ctx.db.scalars(_q_notas().where(*f).order_by(Nfse.numero.desc()).offset((pagina - 1) * por_pagina).limit(por_pagina))
    return Pagina[NfseOut](itens=[_nfse_out(n) for n in itens], total=total, pagina=pagina, por_pagina=por_pagina)


def _nota(ctx, nid: int) -> Nfse:
    return _do_tenant(ctx, Nfse, nid)


@router.get("/notas/email/modelo", response_model=ModeloEmailOut)
def modelo_email(ctx: Ctx):
    """Modelo com variáveis para o envio em lote (cada nota recebe o seu texto)."""
    from app.core.models import ConfigNotificacao

    cfg = ctx.db.get(ConfigNotificacao, ctx.tenant_id)
    nfse_cfg = service.config_nfse(ctx.db, ctx.tenant_id)
    padrao = [t for t, ligado in (("pdf", not cfg or cfg.anexar_pdf), ("xml", not cfg or cfg.anexar_xml),
                                  ("fatura", nfse_cfg.fatura_anexar_email)) if ligado]
    return ModeloEmailOut(
        remetente=f"{cfg.remetente_nome} <{cfg.remetente_email}>" if cfg else settings.smtp_from,
        assunto_modelo=cfg.assunto_modelo if cfg else "NFS-e nº {numero_nota}",
        corpo_modelo=cfg.corpo_modelo if cfg else "Segue a NFS-e nº {numero_nota}.",
        assinatura=service.assinatura_email(ctx.db, ctx.tenant_id),
        cc=list(cfg.email_cc_padrao or []) if cfg else [], cco=list(cfg.email_cco_padrao or []) if cfg else [],
        anexos_padrao=padrao, variaveis=service.VARIAVEIS_EMAIL,
    )


@router.post("/notas/enviar-emails", response_model=EnvioLoteOut)
def enviar_emails_selecionadas(dados: EnviarEmailsLoteIn, ctx: Operador):
    """Envia as notas selecionadas (ou todas do filtro), cada uma para o próprio cliente."""
    if not set(dados.anexos) <= {"pdf", "xml", "fatura"}:
        raise HTTPException(400, "Anexo desconhecido.")
    if dados.ids:
        notas = list(ctx.db.scalars(select(Nfse).where(Nfse.tenant_id == ctx.tenant_id, Nfse.id.in_(dados.ids))))
    elif dados.filtro:
        fl = dados.filtro
        f = _filtro_notas(ctx, fl.q, fl.status, fl.emissora_id, fl.lote_id, fl.de, fl.ate, None, None, fl.pagamento)
        notas = list(ctx.db.scalars(_q_notas().where(*f).order_by(Nfse.numero)))
    else:
        raise HTTPException(400, "Selecione as notas ou use \"todas do filtro\".")
    notas = [n for n in notas if n.status == StatusNfse.emitida]
    if len(notas) > 1000:
        raise HTTPException(400, "No máximo 1.000 notas por envio. Refine o filtro.")
    assinatura = service.assinatura_email(ctx.db, ctx.tenant_id)
    enviados, puladas, falhas, sem_email = 0, 0, [], []
    for n in notas:
        if dados.pular_ja_enviadas and any(e.status == "enviado" for e in n.envios):
            puladas += 1
            continue
        para = list(dict.fromkeys(service.emails_do_cliente(n) + [str(e).lower() for e in dados.para_extra]))
        if not service.emails_do_cliente(n):
            sem_email.append({"nfse_id": n.id, "numero": n.numero,
                              "cliente": (n.rps.dados_json or {}).get("tomador", {}).get("nome", "")})
            continue
        v = service.variaveis_nota(ctx.db, n)
        corpo = service.aplicar_modelo(dados.corpo_modelo, v) + (f"\n\n{assinatura}" if assinatura else "")
        envio = service.enviar_email_nota(
            ctx.db, n, para, cc=[str(e).lower() for e in dados.cc], cco=[str(e).lower() for e in dados.cco],
            assunto=service.aplicar_modelo(dados.assunto_modelo, v), corpo=corpo, anexos=dados.anexos,
            usuario_nome=ctx.usuario.nome,
        )
        if envio.status == "enviado":
            enviados += 1
        else:
            falhas.append({"nfse_id": n.id, "numero": n.numero, "erro": envio.erro})
    ctx.auditar("nfse.enviar_emails_lote", "nfse",
                descricao=f"{enviados} enviados, {len(falhas)} falhas, {len(sem_email)} sem e-mail, {puladas} puladas")
    ctx.db.commit()
    return EnvioLoteOut(total=len(notas), enviados=enviados, falhas=falhas, sem_email=sem_email, puladas=puladas)


@router.get("/notas/{nid}", response_model=NfseDetalhe)
def ler_nota(nid: int, ctx: Ctx):
    n = _nota(ctx, nid)
    tom = (n.rps.dados_json or {}).get("tomador", {})
    emails = (n.tomador.emails if n.tomador else []) or ([tom["email"]] if tom.get("email") else [])
    return NfseDetalhe(
        **_nfse_out(n).model_dump(), discriminacao=n.rps.discriminacao, cod_servico=n.rps.cod_servico,
        aliquota=n.rps.aliquota, tomador_emails=emails,
        envios=[_envio_out(e) for e in sorted(n.envios, key=lambda e: e.enviado_em, reverse=True)],
    )


@router.get("/notas/{nid}/pdf")
def pdf_nota(nid: int, ctx: Ctx):
    n = _nota(ctx, nid)
    if not n.pdf_path or not Path(n.pdf_path).exists():
        raise HTTPException(404, "PDF não encontrado.")
    return FileResponse(n.pdf_path, media_type="application/pdf", filename=f"NFSe_{n.numero}.pdf",
                        content_disposition_type="inline")


@router.get("/notas/{nid}/xml")
def xml_nota(nid: int, ctx: Ctx):
    n = _nota(ctx, nid)
    if not n.xml_path or not Path(n.xml_path).exists():
        raise HTTPException(404, "XML não encontrado.")
    return FileResponse(n.xml_path, media_type="application/xml", filename=f"NFSe_{n.numero}.xml")


def _notas_selecionadas(ctx, dados: ZipIn) -> list[Nfse]:
    if dados.lote_id:
        _do_tenant(ctx, Lote, dados.lote_id)
        return list(ctx.db.scalars(_q_notas().where(Nfse.tenant_id == ctx.tenant_id, Rps.lote_id == dados.lote_id)))
    if not dados.ids:
        raise HTTPException(400, "Selecione ao menos uma nota.")
    return list(ctx.db.scalars(select(Nfse).where(Nfse.tenant_id == ctx.tenant_id, Nfse.id.in_(dados.ids))))


@router.post("/notas/zip")
def zip_notas(dados: ZipIn, ctx: Ctx):
    notas = _notas_selecionadas(ctx, dados)
    nome = f"notas_lote_{dados.lote_id}.zip" if dados.lote_id else f"notas_{date.today():%Y%m%d}.zip"
    return _baixar(service.zip_notas(notas), nome, "application/zip")


@router.post("/notas/exportar")
def exportar_notas(dados: ZipIn, ctx: Ctx):
    notas = _notas_selecionadas(ctx, dados)
    wb = Workbook()
    ws = wb.active
    ws.title = "NFS-e"
    ws.append(["Nº NFS-e", "Cód. verificação", "Emissão", "RPS", "Tomador", "CPF/CNPJ", "Valor", "ISS", "Status", "Link"])
    for n in notas:
        o = _nfse_out(n)
        ws.append([o.numero, o.cod_verificacao, o.data_emissao.replace(tzinfo=None), f"{o.rps_serie}-{o.rps_numero}",
                   o.tomador_nome, o.tomador_doc, float(o.valor), float(o.iss), o.status.value, o.link_prefeitura])
    buf = io.BytesIO()
    wb.save(buf)
    return _baixar(buf.getvalue(), f"notas_{date.today():%Y%m%d}.xlsx", XLSX)


@router.post("/notas/{nid}/cancelar", response_model=NfseOut)
def cancelar(nid: int, dados: CancelarIn, ctx: Admin):
    n = _nota(ctx, nid)
    if n.status == StatusNfse.cancelada:
        raise HTTPException(400, "Esta nota já está cancelada.")
    if not n.simulada:
        cert_reg = ctx.db.get(Certificado, n.emissora.certificado_id) if n.emissora.certificado_id else None
        if not cert_reg:
            raise HTTPException(400, "Sem certificado digital ativo: cancele pelo portal da Prefeitura.")
        try:
            modo = service.modo_prefeitura(ctx.db, ctx.tenant_id)
            ret = ClientePrefeitura(certificados.abrir(cert_reg), modo).cancelar(cnpj=n.emissora.cnpj, ccm=n.emissora.ccm, numero_nfe=n.numero)
        except (FalhaComunicacao, RuntimeError) as e:
            raise HTTPException(502, f"A Prefeitura não respondeu. Tente de novo em instantes. ({e})") from e
        if not ret.sucesso:
            h = ret.erros[0].humano() if ret.erros else {"mensagem": "A Prefeitura recusou o cancelamento.", "acao": ""}
            raise HTTPException(400, f"{h['mensagem']} {h.get('acao', '')}".strip())
    n.status = StatusNfse.cancelada
    n.cancelada_em = agora()
    n.motivo_cancelamento = dados.motivo
    fatura.sincronizar(ctx.db, n)
    ctx.auditar("nfse.cancelar", "nfse", n.id, descricao=f"NFS-e {n.numero} cancelada: {dados.motivo}",
                antes={"status": "emitida"}, depois={"status": "cancelada", "motivo": dados.motivo})
    ctx.db.commit()
    return _nfse_out(n)


def _envio_out(e) -> dict:
    return {"destinatarios": e.destinatarios, "cc": e.cc or [], "cco": e.cco or [], "assunto": e.assunto,
            "anexos": e.anexos or [], "usuario_nome": e.usuario_nome, "status": e.status,
            "enviado_em": e.enviado_em.isoformat(), "erro": e.erro}


@router.get("/notas/{nid}/email/previa", response_model=EmailPreviaOut)
def previa_email(nid: int, ctx: Ctx):
    """O que o compositor de e-mail mostra: modelo já preenchido, cópias padrão e anexos disponíveis."""
    n = _nota(ctx, nid)
    m = service.montar_email(ctx.db, n)
    return EmailPreviaOut(
        remetente=m["remetente"], para=m["para"], cc=m["cc"], cco=m["cco"], assunto=m["assunto"], corpo=m["corpo"],
        anexos=[AnexoEmail(tipo=a["tipo"], nome=a["nome"], tamanho=a["tamanho"], padrao=a["padrao"]) for a in m["anexos"]],
        emails_cliente=m["emails_cliente"], cliente_id=n.tomador_id,
        envios=[_envio_out(e) for e in sorted(n.envios, key=lambda e: e.enviado_em, reverse=True)],
    )


@router.post("/notas/{nid}/enviar-email", response_model=Mensagem)
def enviar_email(nid: int, dados: EnviarEmailIn, ctx: Operador):
    n = _nota(ctx, nid)
    if n.status == StatusNfse.cancelada:
        raise HTTPException(400, "Nota cancelada não é enviada ao cliente.")
    if dados.anexos is not None and not set(dados.anexos) <= {"pdf", "xml", "fatura"}:
        raise HTTPException(400, "Anexo desconhecido.")
    para = [str(e).lower() for e in dados.destinatarios] if dados.destinatarios else None
    try:
        envio = service.enviar_email_nota(
            ctx.db, n, para, cc=[str(e).lower() for e in dados.cc] if dados.cc is not None else None,
            cco=[str(e).lower() for e in dados.cco] if dados.cco is not None else None,
            assunto=dados.assunto, corpo=dados.corpo, anexos=dados.anexos, usuario_nome=ctx.usuario.nome,
        )
    except service.RegraNegocio as e:
        raise _erro(e) from e
    novos = []
    if dados.salvar_no_cliente and para and n.tomador:
        novos = [e for e in para if e not in (n.tomador.emails or [])]
        if novos:
            n.tomador.emails = [*(n.tomador.emails or []), *novos]
    ctx.auditar("nfse.enviar_email", "nfse", n.id,
                descricao=f"NFS-e {n.numero} → {', '.join(envio.destinatarios)}"
                + (f" · Cc {', '.join(envio.cc)}" if envio.cc else "") + (f" · Cco {len(envio.cco)}" if envio.cco else ""))
    ctx.db.commit()
    if envio.status != "enviado":
        raise HTTPException(502, f"Falha no envio: {envio.erro}")
    extra = f" · {len(novos)} e-mail(s) novo(s) salvo(s) no cadastro do cliente" if novos else ""
    return Mensagem(mensagem=f"Nota enviada para {', '.join(envio.destinatarios)}.{extra}")



# ======================================================================== painel
def _ind(atual: float, anterior: float) -> Indicador:
    var = None if not anterior else round((atual - anterior) / anterior * 100, 1)
    return Indicador(valor=atual, anterior=anterior, variacao=var)


@router.get("/painel", response_model=PainelOut)
def painel(ctx: Ctx):
    db = ctx.db
    hoje = date.today()
    ini = hoje.replace(day=1)
    ini_ant = (ini - timedelta(days=1)).replace(day=1)

    def periodo(a: date, b: date):
        base = [Nfse.tenant_id == ctx.tenant_id, Nfse.data_emissao >= datetime.combine(a, time.min),
                Nfse.data_emissao < datetime.combine(b, time.min)]
        n, v, i = db.execute(
            select(func.count(Nfse.id), func.coalesce(func.sum(Nfse.valor), 0), func.coalesce(func.sum(Nfse.iss), 0))
            .where(*base, Nfse.status == StatusNfse.emitida)
        ).one()
        lotes = list(db.scalars(select(Lote).where(Lote.tenant_id == ctx.tenant_id, Lote.concluido_em.is_not(None),
                                                   Lote.concluido_em >= datetime.combine(a, time.min),
                                                   Lote.concluido_em < datetime.combine(b, time.min))))
        rej = sum(lt.qtd_rejeitadas for lt in lotes)
        tempos = [(lt.concluido_em - lt.iniciado_em).total_seconds() for lt in lotes if lt.iniciado_em]
        return int(n), float(v), float(i), rej, (sum(tempos) / len(tempos) if tempos else 0)

    prox = (ini + timedelta(days=32)).replace(day=1)
    atual, ant = periodo(ini, prox), periodo(ini_ant, ini)

    por_dia_q = db.execute(
        select(func.date(Nfse.data_emissao), func.count(Nfse.id), func.sum(Nfse.valor))
        .where(Nfse.tenant_id == ctx.tenant_id, Nfse.status == StatusNfse.emitida,
               Nfse.data_emissao >= datetime.combine(hoje - timedelta(days=29), time.min))
        .group_by(func.date(Nfse.data_emissao))
    ).all()
    mapa = {str(d): (n, float(v or 0)) for d, n, v in por_dia_q}
    por_dia = []
    for k in range(29, -1, -1):
        d = (hoje - timedelta(days=k)).isoformat()
        n, v = mapa.get(d, (0, 0.0))
        por_dia.append({"dia": d, "notas": n, "valor": v})

    top_q = db.execute(
        select(Tomador.nome, func.sum(Nfse.valor), func.count(Nfse.id))
        .join(Tomador, Nfse.tomador_id == Tomador.id)
        .where(Nfse.tenant_id == ctx.tenant_id, Nfse.status == StatusNfse.emitida,
               Nfse.data_emissao >= datetime.combine(ini, time.min))
        .group_by(Tomador.nome).order_by(func.sum(Nfse.valor).desc()).limit(10)
    ).all()

    status_q = db.execute(
        select(Lote.status, func.count(Lote.id))
        .where(Lote.tenant_id == ctx.tenant_id, Lote.status.in_([StatusLote.concluido, StatusLote.parcial, StatusLote.falhou]))
        .group_by(Lote.status)
    ).all()

    alertas: list[Alerta] = []
    cert = db.scalar(select(Certificado).where(Certificado.tenant_id == ctx.tenant_id, Certificado.ativo))
    if cert:
        dias = (cert.valido_ate - hoje).days
        if dias < 0:
            alertas.append(Alerta(nivel="erro", texto=f"Seu certificado digital venceu há {-dias} dias.", href="/configuracoes/certificado"))
        elif dias <= 30:
            alertas.append(Alerta(nivel="aviso", texto=f"Seu certificado digital vence em {dias} dias.", href="/configuracoes/certificado"))
    elif service.modo_prefeitura(db, ctx.tenant_id) != "simulado":
        alertas.append(Alerta(nivel="erro", texto="Nenhum certificado digital ativo. Sem ele não há emissão.", href="/configuracoes/certificado"))
    for lt in db.scalars(select(Lote).where(Lote.tenant_id == ctx.tenant_id, Lote.qtd_rejeitadas > 0,
                                            Lote.status.in_([StatusLote.parcial, StatusLote.falhou])).order_by(Lote.id.desc()).limit(3)):
        alertas.append(Alerta(nivel="erro", texto=f"{lt.qtd_rejeitadas} nota(s) do lote #{lt.numero:04d} foram rejeitadas — ver motivo.",
                              href=f"/nfse/lotes/{lt.id}"))
    fila = db.scalar(select(func.count(Lote.id)).where(Lote.tenant_id == ctx.tenant_id, Lote.status.in_([StatusLote.na_fila, StatusLote.transmitindo])))
    if fila:
        alertas.append(Alerta(nivel="info", texto=f"{fila} lote(s) sendo transmitidos agora; acompanhe ao vivo.", href="/nfse/lotes"))
    if service.modo_prefeitura(db, ctx.tenant_id) == "simulado":
        alertas.append(Alerta(nivel="info", texto="Modo simulação: as notas são de demonstração e não têm valor fiscal."))

    a_receber, a_receber_qtd = db.execute(
        select(func.coalesce(func.sum(Nfse.valor), 0), func.count(Nfse.id)).where(
            Nfse.tenant_id == ctx.tenant_id, Nfse.status == StatusNfse.emitida, Nfse.pago_em.is_(None))
    ).one()
    recebido_mes = db.scalar(
        select(func.coalesce(func.sum(Nfse.valor_pago), 0)).where(Nfse.tenant_id == ctx.tenant_id, Nfse.pago_em >= ini)
    ) or 0
    ultimos = db.scalars(select(Lote).where(Lote.tenant_id == ctx.tenant_id, Lote.status != StatusLote.descartado).order_by(Lote.id.desc()).limit(5))
    return PainelOut(
        a_receber=float(a_receber), a_receber_qtd=int(a_receber_qtd), recebido_mes=float(recebido_mes),
        notas_emitidas=_ind(atual[0], ant[0]), valor_faturado=_ind(atual[1], ant[1]), iss=_ind(atual[2], ant[2]),
        rejeicoes=_ind(atual[3], ant[3]), tempo_medio_lote_seg=_ind(round(atual[4], 1), round(ant[4], 1)),
        por_dia=por_dia, top_tomadores=[{"nome": n, "valor": float(v), "notas": c} for n, v, c in top_q],
        status_lotes=[{"status": s.value if hasattr(s, "value") else s, "qtd": q} for s, q in status_q],
        alertas=alertas, ultimos_lotes=[_lote_out(lt) for lt in ultimos],
    )

