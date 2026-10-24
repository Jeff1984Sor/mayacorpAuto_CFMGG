"""Leitura de extrato bancário: OFX (padrão de todos os bancos), CSV e Excel.

Só os créditos interessam: são os pagamentos que a Lello fez à CFME.
"""

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from app.modules.rpa_nfse.parser_txt import decodificar


class ExtratoInvalido(ValueError):
    pass


@dataclass
class Credito:
    data: date
    valor: Decimal
    descricao: str
    documento: str | None
    chave: str


@dataclass
class ExtratoLido:
    formato: str
    banco: str | None
    creditos: list[Credito]
    debitos_ignorados: int


def _norm(txt: str) -> str:
    return unicodedata.normalize("NFKD", txt or "").encode("ascii", "ignore").decode().lower().strip()


def _chave(data: date, valor: Decimal, descricao: str, doc: str | None) -> str:
    base = f"{data.isoformat()}|{valor}|{_norm(descricao)}|{doc or ''}"
    return "h:" + hashlib.sha1(base.encode()).hexdigest()[:32]  # noqa: S324 (identidade, não segurança)


def _valor_br(txt) -> Decimal | None:
    if txt is None or txt == "":
        return None
    if isinstance(txt, int | float | Decimal):
        return Decimal(str(txt)).quantize(Decimal("0.01"))
    s = str(txt).strip().replace("R$", "").replace(" ", "")
    negativo = s.endswith("-") or s.startswith("-") or (s.startswith("(") and s.endswith(")")) or s.upper().endswith("D")
    s = s.strip("-()").rstrip("CDcd")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        v = Decimal(s).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None
    return -v if negativo else v


def _data(txt) -> date | None:
    if isinstance(txt, datetime):
        return txt.date()
    if isinstance(txt, date):
        return txt
    s = str(txt or "").strip()[:10]
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%Y", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


# ---------------------------------------------------------------- OFX
def _ler_ofx(texto: str) -> ExtratoLido:
    banco = None
    if m := re.search(r"<ORG>([^<\r\n]+)", texto):
        banco = m.group(1).strip()
    creditos, debitos = [], 0
    for bloco in re.findall(r"<STMTTRN>(.*?)(?:</STMTTRN>|(?=<STMTTRN>)|(?=</BANKTRANLIST>))", texto, re.S | re.I):
        def tag(nome: str, bloco: str = bloco) -> str | None:
            m = re.search(rf"<{nome}>([^<\r\n]*)", bloco, re.I)
            return m.group(1).strip() if m else None

        valor = _valor_br((tag("TRNAMT") or "").replace(",", "."))
        quando = _data((tag("DTPOSTED") or "")[:8])
        if valor is None or quando is None:
            continue
        if valor <= 0:
            debitos += 1
            continue
        descricao = " ".join(filter(None, [tag("NAME"), tag("MEMO")]))
        fitid = tag("FITID")
        creditos.append(Credito(quando, valor, descricao, tag("CHECKNUM") or tag("REFNUM"),
                                f"ofx:{fitid}" if fitid else _chave(quando, valor, descricao, None)))
    return ExtratoLido("ofx", banco, creditos, debitos)


# ---------------------------------------------------------------- CSV / Excel
COLUNAS = {
    "data": ("data", "dt", "data lancamento", "data do lancamento", "data mov", "data movimento"),
    "descricao": ("descricao", "historico", "lancamento", "memo", "detalhe", "descricao do lancamento"),
    "valor": ("valor", "valor (r$)", "valor r$", "montante"),
    "credito": ("credito", "entrada", "credito (r$)"),
    "debito": ("debito", "saida", "debito (r$)"),
    "documento": ("documento", "doc", "n documento", "numero documento", "num doc"),
}


def _mapear(cabecalho: list) -> dict[str, int]:
    norm = [_norm(str(c or "")) for c in cabecalho]
    mapa = {}
    for campo, nomes in COLUNAS.items():
        for i, c in enumerate(norm):
            if c in nomes or any(c.startswith(n) for n in nomes):
                mapa.setdefault(campo, i)
    return mapa


def _ler_tabela(linhas: list[list], formato: str) -> ExtratoLido:
    inicio, mapa = None, {}
    for i, ln in enumerate(linhas[:30]):  # bancos colocam algumas linhas de título antes do cabeçalho
        m = _mapear(ln)
        if "data" in m and ("valor" in m or "credito" in m):
            inicio, mapa = i, m
            break
    if inicio is None:
        raise ExtratoInvalido("Não encontrei as colunas de data e valor. O extrato precisa ter cabeçalho (Data, Histórico, Valor).")
    creditos, debitos = [], 0
    for ln in linhas[inicio + 1:]:
        def col(campo, ln=ln):
            i = mapa.get(campo)
            return ln[i] if i is not None and i < len(ln) else None

        quando = _data(col("data"))
        if not quando:
            continue
        valor = _valor_br(col("credito")) if "credito" in mapa else _valor_br(col("valor"))
        if "credito" in mapa and not valor and "debito" in mapa and _valor_br(col("debito")):
            debitos += 1
            continue
        if valor is None or valor <= 0:
            debitos += 1 if valor else 0
            continue
        desc = str(col("descricao") or "").strip()
        doc = str(col("documento") or "").strip() or None
        creditos.append(Credito(quando, valor, desc, doc, _chave(quando, valor, desc, doc)))
    return ExtratoLido(formato, None, creditos, debitos)


def ler(nome: str, conteudo: bytes) -> ExtratoLido:
    ext = nome.lower().rsplit(".", 1)[-1] if "." in nome else ""
    if ext in ("xlsx", "xlsm"):
        ws = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).active
        lido = _ler_tabela([list(r) for r in ws.iter_rows(values_only=True)], "xlsx")
    else:
        texto, _ = decodificar(conteudo)
        if ext == "ofx" or "<OFX>" in texto.upper() or "OFXHEADER" in texto.upper():
            lido = _ler_ofx(texto)
        elif ext in ("csv", "txt"):
            sep = ";" if texto.count(";") >= texto.count(",") else ","
            lido = _ler_tabela(list(csv.reader(io.StringIO(texto), delimiter=sep)), "csv")
        else:
            raise ExtratoInvalido("Envie o extrato em OFX (recomendado), CSV ou Excel.")
    if not lido.creditos:
        raise ExtratoInvalido("Nenhum crédito encontrado no extrato.")
    # o mesmo crédito repetido no arquivo (FITIDs iguais) conta uma vez
    vistos, unicos = set(), []
    for c in lido.creditos:
        if c.chave not in vistos:
            vistos.add(c.chave)
            unicos.append(c)
    lido.creditos = sorted(unicos, key=lambda c: c.data)
    return lido
