"""Configuração da integração: escritório (modo, notificações), plataforma MayaCorp (tribunais, sistemas, tabelas, planos)
e cobrança. Inclui o link público do recibo (token com validade)."""

import csv
import io
from decimal import Decimal
from pathlib import Path

import jwt
from fastapi import APIRouter, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse
from sqlalchemy import func, or_, select

from app.core.config import settings
from app.core.deps import DB, Ctx, EquipeMayaCorp
from app.modules.tribunais import catalogo, cobranca, jobs, notificacoes, service
from app.modules.tribunais.models import (
    Consumo,
    FaixaExcedente,
    FaixaSetup,
    NotificacaoProtocolo,
    PesoProtocolo,
    Plano,
    RetornoProtocolo,
    Sistema,
    TabelaLocal,
    Tribunal,
)
from app.modules.tribunais.routes import _dt
from app.modules.tribunais.schemas import (
    AjusteIn,
    ConfigIn,
    EstornoIn,
    FaixaSetupIn,
    NotificacaoIn,
    PlanoIn,
    PlataformaIn,
    SistemaIn,
)

router = APIRouter(prefix="/tribunais", tags=["Tribunais · Configuração"])
publico = APIRouter(prefix="/public/tribunais", tags=["público"])


# ---------------------------------------------------------------------------------------------------- escritório
def _config_out(ctx) -> dict:
    cfg = service.config(ctx.db, ctx.tenant_id)
    plano = ctx.db.get(Plano, cfg.plano_id) if cfg.plano_id else None
    return {"modo": cfg.modo, "sincronizar_auto": cfg.sincronizar_auto, "criar_andamentos": cfg.criar_andamentos,
            "vincular_automatico": cfg.vincular_automatico, "limiar_fuzzy": cfg.limiar_fuzzy, "plano_id": cfg.plano_id,
            "plano_nome": plano.nome if plano else None, "ambiente": settings.ambiente,
            "modo_real_disponivel": settings.ambiente == "prod"}


@router.get("/config")
def ler_config(ctx: Ctx):
    out = _config_out(ctx)
    ctx.db.commit()
    return out


@router.put("/config")
def salvar_config(dados: ConfigIn, ctx: Ctx):
    if dados.modo == "real" and settings.ambiente != "prod":
        raise HTTPException(400, "O modo real só existe em produção. Aqui (homologação/dev) o protocolo é sempre simulado.")
    cfg = service.config(ctx.db, ctx.tenant_id)
    antes = cfg.modo
    for k, v in dados.model_dump().items():
        setattr(cfg, k, v)
    ctx.auditar("tribunais.config.alterar", "trb_config", ctx.tenant_id, descricao=f"modo {antes} → {cfg.modo}" if antes != cfg.modo else None,
                depois=dados.model_dump())
    ctx.db.commit()
    return _config_out(ctx)


def _notif_out(n: NotificacaoProtocolo) -> dict:
    return {"evento": n.evento, "ativo": n.ativo, "notificar_responsavel_caso": n.notificar_responsavel_caso,
            "notificar_quem_protocolou": n.notificar_quem_protocolou, "notificar_titular_certificado": n.notificar_titular_certificado,
            "destinatarios_fixos": n.destinatarios_fixos, "copia_oculta": n.copia_oculta, "anexar_recibo": n.anexar_recibo,
            "anexar_peca_principal": n.anexar_peca_principal, "respeitar_sigilo": n.respeitar_sigilo, "assunto": n.assunto,
            "corpo_texto": n.corpo_texto, "corpo_html": n.corpo_html}


VARIAVEIS = ["advogado_responsavel_nome", "protocolado_por_nome", "cliente_nome", "numero_cnj", "numero_protocolo", "data_hora_protocolo",
             "tribunal_sigla", "sistema", "orgao_julgador", "classe", "tipo_peticao", "partes_resumo", "valor_causa", "prazo_final",
             "mensagem_erro", "link_processo"]


@router.get("/notificacoes")
def notificacoes_cfg(ctx: Ctx):
    service.config(ctx.db, ctx.tenant_id)
    lista = ctx.db.scalars(select(NotificacaoProtocolo).where(NotificacaoProtocolo.tenant_id == ctx.tenant_id))
    ordem = {e: i for i, e in enumerate(catalogo.EVENTOS)}
    out = {"eventos": sorted((_notif_out(n) for n in lista), key=lambda n: ordem.get(n["evento"], 9)), "variaveis": VARIAVEIS}
    ctx.db.commit()
    return out


@router.put("/notificacoes/{evento}")
def salvar_notificacao(evento: str, dados: NotificacaoIn, ctx: Ctx):
    service.config(ctx.db, ctx.tenant_id)
    n = ctx.db.scalar(select(NotificacaoProtocolo).where(NotificacaoProtocolo.tenant_id == ctx.tenant_id, NotificacaoProtocolo.evento == evento))
    if not n:
        raise HTTPException(404, "Evento desconhecido.")
    for k, v in dados.model_dump().items():
        setattr(n, k, [str(x) for x in v] if k in ("destinatarios_fixos", "copia_oculta") else v)
    ctx.auditar("tribunais.notificacao.alterar", "trb_notif_config", n.id, descricao=evento)
    ctx.db.commit()
    return _notif_out(n)


@router.post("/notificacoes/{evento}/restaurar")
def restaurar_notificacao(evento: str, ctx: Ctx):
    n = ctx.db.scalar(select(NotificacaoProtocolo).where(NotificacaoProtocolo.tenant_id == ctx.tenant_id, NotificacaoProtocolo.evento == evento))
    if not n or evento not in catalogo.MODELOS:
        raise HTTPException(404, "Evento desconhecido.")
    n.assunto, n.corpo_texto, n.corpo_html = catalogo.MODELOS[evento]["assunto"], catalogo.MODELOS[evento]["corpo_texto"], None
    ctx.db.commit()
    return _notif_out(n)


# ------------------------------------------------------------------------------------------- tribunais e sistemas
def sistema_out(s: Sistema) -> dict:
    return {c: (getattr(s, c).isoformat() if hasattr(getattr(s, c), "isoformat") else getattr(s, c)) for c in (
        "id", "tribunal_id", "grau", "sistema", "versao_sistema", "url_portal", "url_consulta_publica", "url_mni", "versao_mni",
        "mni_consulta", "mni_peticionamento", "mni_inicial", "exige_credenciamento", "autenticacao", "tamanho_max_arquivo_mb",
        "tamanho_max_total_mb", "formatos_aceitos", "exige_pdf_pesquisavel", "formato_assinatura", "adaptador", "config_extra",
        "data_inicio", "data_fim", "status_saude", "ultima_verificacao", "validado", "intervalo_min_s", "janela_inicio", "janela_fim",
        "ativo")}


@router.get("/tribunais")
def tribunais(ctx: Ctx, ramo: str | None = None, q: str | None = None):
    service.plataforma(ctx.db)
    sel = select(Tribunal)
    if ramo:
        sel = sel.where(Tribunal.ramo == ramo)
    if q:
        sel = sel.where(or_(Tribunal.sigla.ilike(f"%{q}%"), Tribunal.nome.ilike(f"%{q}%")))
    ordem = {"superior": 0, "estadual": 1, "federal": 2, "trabalho": 3, "eleitoral": 4, "militar": 5}
    lista = sorted(ctx.db.scalars(sel), key=lambda t: (ordem.get(t.ramo, 9), t.segmento_j, t.codigo_tr))
    out = [{"id": t.id, "sigla": t.sigla, "nome": t.nome, "segmento_j": t.segmento_j, "codigo_tr": t.codigo_tr, "uf": t.uf, "ramo": t.ramo,
            "alias_datajud": t.alias_datajud, "sistemas": [sistema_out(s) for s in t.sistemas]} for t in lista]
    ctx.db.commit()
    return out


@router.put("/sistemas/{sid}")
def salvar_sistema(sid: int, dados: SistemaIn, ctx: EquipeMayaCorp):
    s = ctx.db.get(Sistema, sid)
    if not s:
        raise HTTPException(404, "Sistema não encontrado.")
    if dados.validado and not (dados.url_mni or dados.formato_assinatura == "assinatura_no_portal"):
        raise HTTPException(400, "Para homologar o protocolo automático, informe o endereço do MNI.")
    for k, v in dados.model_dump().items():
        setattr(s, k, v)
    ctx.auditar("tribunais.sistema.alterar", "trb_sistema", s.id, descricao=f"{s.tribunal.sigla} {s.grau} {s.sistema}", depois=dados.model_dump(mode="json"))
    ctx.db.commit()
    return sistema_out(s)


@router.post("/tribunais/{tid}/sistemas", status_code=201)
def novo_sistema(tid: int, dados: SistemaIn, ctx: EquipeMayaCorp):
    t = ctx.db.get(Tribunal, tid)
    if not t:
        raise HTTPException(404, "Tribunal não encontrado.")
    s = Sistema(tribunal_id=t.id, **dados.model_dump())
    ctx.db.add(s)
    ctx.db.flush()
    ctx.auditar("tribunais.sistema.criar", "trb_sistema", s.id, descricao=f"{t.sigla} {s.grau} {s.sistema}")
    ctx.db.commit()
    return sistema_out(s)


@router.delete("/sistemas/{sid}")
def excluir_sistema(sid: int, ctx: EquipeMayaCorp):
    from app.modules.tribunais.models import Processo, Protocolo

    s = ctx.db.get(Sistema, sid)
    if not s:
        raise HTTPException(404, "Sistema não encontrado.")
    em_uso = ctx.db.scalar(select(Processo.id).where(Processo.sistema_instancia_id == s.id)) or ctx.db.scalar(
        select(Protocolo.id).where(Protocolo.sistema_instancia_id == s.id))
    if em_uso:
        s.ativo = False  # tem histórico: só desativa
    else:
        ctx.db.delete(s)
    ctx.auditar("tribunais.sistema.excluir", "trb_sistema", sid)
    ctx.db.commit()
    return {"ok": True, "desativado": bool(em_uso)}


@router.post("/sistemas/{sid}/saude")
def checar_saude(sid: int, ctx: EquipeMayaCorp):
    from app.modules.tribunais import adaptadores

    s = ctx.db.get(Sistema, sid)
    if not s:
        raise HTTPException(404, "Sistema não encontrado.")
    from app.core.db import agora

    s.status_saude, s.ultima_verificacao = adaptadores.para(s).saude(s), agora()
    ctx.db.commit()
    return sistema_out(s)


# ---------------------------------------------------------------------------------------------- tabelas locais
@router.get("/tabelas")
def tabelas(ctx: Ctx, tabela: str, tribunal: str | None = None, q: str | None = None, limite: int = Query(50, le=500)):
    sel = select(TabelaLocal).where(TabelaLocal.tabela == tabela, TabelaLocal.ativo.is_(True))
    if tribunal:
        t = ctx.db.scalar(select(Tribunal).where(Tribunal.sigla == tribunal.upper()))
        sel = sel.where(or_(TabelaLocal.tribunal_id.is_(None), TabelaLocal.tribunal_id == (t.id if t else 0)))
    if q:
        sel = sel.where(or_(TabelaLocal.nome.ilike(f"%{q}%"), TabelaLocal.codigo_local == q.strip(), TabelaLocal.codigo_cnj == q.strip()))
    return [{"id": x.id, "tabela": x.tabela, "codigo_local": x.codigo_local, "codigo_cnj": x.codigo_cnj, "nome": x.nome,
             "pai_codigo": x.pai_codigo, "extra": x.extra, "tribunal_id": x.tribunal_id, "sistema_id": x.sistema_id}
            for x in ctx.db.scalars(sel.order_by(TabelaLocal.nome).limit(limite))]


@router.get("/tabelas/resumo")
def tabelas_resumo(ctx: Ctx):
    linhas = ctx.db.execute(select(TabelaLocal.tribunal_id, TabelaLocal.tabela, func.count()).group_by(TabelaLocal.tribunal_id, TabelaLocal.tabela)).all()
    siglas = {t.id: t.sigla for t in ctx.db.scalars(select(Tribunal))}
    return [{"tribunal": siglas.get(t, "Nacional (TPU)") if t else "Nacional (TPU)", "tabela": tab, "registros": n} for t, tab, n in linhas]


@router.post("/tabelas/importar")
async def importar_tabela(ctx: EquipeMayaCorp, tabela: str = Form(...), arquivo: UploadFile = File(...), tribunal: str | None = Form(None),
                          substituir: bool = Form(True)):
    """CSV/XLSX com colunas codigo_local, nome e opcionais codigo_cnj, pai_codigo (+ qualquer outra vai para `extra`)."""
    if tabela not in ("classe", "assunto", "tipo_documento", "tipo_peticao", "localidade", "orgao_julgador", "competencia"):
        raise HTTPException(400, "Tabela inválida.")
    t = ctx.db.scalar(select(Tribunal).where(Tribunal.sigla == tribunal.upper())) if tribunal else None
    if tribunal and not t:
        raise HTTPException(404, "Tribunal não encontrado.")
    dados = await arquivo.read()
    nome = (arquivo.filename or "").lower()
    if nome.endswith(".xlsx"):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(dados), read_only=True, data_only=True).active
        linhas = list(ws.iter_rows(values_only=True))
        cab = [str(c or "").strip().lower() for c in linhas[0]]
        registros = [dict(zip(cab, ["" if v is None else str(v).strip() for v in ln], strict=False)) for ln in linhas[1:]]
    else:
        texto = dados.decode("utf-8-sig", errors="replace")
        sep = ";" if texto.count(";") > texto.count(",") else ","
        registros = [{(k or "").strip().lower(): (v or "").strip() for k, v in r.items()} for r in csv.DictReader(io.StringIO(texto), delimiter=sep)]
    registros = [r for r in registros if r.get("codigo_local") and r.get("nome")]
    if not registros:
        raise HTTPException(400, "Nenhuma linha com codigo_local e nome.")
    if substituir:
        for x in ctx.db.scalars(select(TabelaLocal).where(TabelaLocal.tabela == tabela, TabelaLocal.tribunal_id == (t.id if t else None))
                                if t else select(TabelaLocal).where(TabelaLocal.tabela == tabela, TabelaLocal.tribunal_id.is_(None))):
            ctx.db.delete(x)
        ctx.db.flush()
    for r in registros:
        extra = {k: v for k, v in r.items() if k not in ("codigo_local", "nome", "codigo_cnj", "pai_codigo") and v}
        ctx.db.add(TabelaLocal(tribunal_id=t.id if t else None, tabela=tabela, codigo_local=r["codigo_local"][:40], nome=r["nome"][:300],
                               codigo_cnj=(r.get("codigo_cnj") or "")[:20] or None, pai_codigo=(r.get("pai_codigo") or "")[:40] or None,
                               extra=extra))
    ctx.auditar("tribunais.tabela.importar", "trb_tabela_local", None, descricao=f"{tabela} · {tribunal or 'TPU'} · {len(registros)} linhas")
    ctx.db.commit()
    return {"importados": len(registros)}


# ------------------------------------------------------------------------------------------------------ plataforma
@router.get("/plataforma")
def ler_plataforma(ctx: EquipeMayaCorp):
    p = service.plataforma(ctx.db)
    ctx.db.commit()
    return {"datajud_url": p.datajud_url, "datajud_chave": p.datajud_chave, "djen_url": p.djen_url, "djen_proxy_url": p.djen_proxy_url,
            "intervalo_datajud_ms": p.intervalo_datajud_ms, "termo_versao": p.termo_versao, "termo_texto": p.termo_texto or catalogo.TERMO_PADRAO}


@router.put("/plataforma")
def salvar_plataforma(dados: PlataformaIn, ctx: EquipeMayaCorp):
    p = service.plataforma(ctx.db)
    for k, v in dados.model_dump().items():
        setattr(p, k, v)
    ctx.auditar("tribunais.plataforma.alterar", "trb_plataforma", p.id)
    ctx.db.commit()
    return ler_plataforma(ctx)


@router.post("/plataforma/testar")
def testar_plataforma(ctx: EquipeMayaCorp):
    """Testa DataJud (pelo servidor) e DJEN (pelo servidor/proxy). O DJEN fora do Brasil cai no navegador."""
    from app.modules.tribunais import fontes

    p = service.plataforma(ctx.db)
    out = {}
    try:
        fontes.datajud_buscar(p.datajud_url, p.datajud_chave, "tjsp", "00000000000000000000", p.intervalo_datajud_ms)
        out["datajud"] = "ok"
    except fontes.FonteIndisponivel as e:
        out["datajud"] = str(e)
    try:
        fontes.djen_buscar(p.djen_proxy_url or p.djen_url, "1", "SP", 1)
        out["djen"] = "ok (servidor)"
    except fontes.ForaDoBrasil:
        out["djen"] = "bloqueado fora do Brasil: a varredura roda pelo navegador" + (" (o proxy também foi recusado)" if p.djen_proxy_url else "")
    except fontes.FonteIndisponivel as e:
        out["djen"] = str(e)
    return out


# --------------------------------------------------------------------------------------------- planos e cobrança
def plano_out(p: Plano) -> dict:
    return {"id": p.id, "nome": p.nome, "preco_mensal": str(p.preco_mensal), "creditos_protocolo_mes": p.creditos_protocolo_mes,
            "creditos_acumulam": p.creditos_acumulam, "permite_excedente": p.permite_excedente, "limite_excedente_mes": p.limite_excedente_mes,
            "limite_processos_monitorados": p.limite_processos_monitorados, "limite_certificados": p.limite_certificados, "ativo": p.ativo,
            "faixas": [{"de_creditos": f.de_creditos, "ate_creditos": f.ate_creditos, "preco_por_credito": str(f.preco_por_credito)}
                       for f in p.faixas]}


@router.get("/planos")
def planos(ctx: EquipeMayaCorp):
    return {"planos": [plano_out(p) for p in ctx.db.scalars(select(Plano).order_by(Plano.preco_mensal))],
            "pesos": {p.tipo_protocolo: p.creditos for p in ctx.db.scalars(select(PesoProtocolo))},
            "setup": [{"id": f.id, "de_processos": f.de_processos, "ate_processos": f.ate_processos, "preco": str(f.preco),
                       "inclui_download_autos": f.inclui_download_autos, "reimportacao_gratis_dias": f.reimportacao_gratis_dias}
                      for f in ctx.db.scalars(select(FaixaSetup).order_by(FaixaSetup.de_processos))]}


def _aplicar_plano(p: Plano, dados: PlanoIn) -> None:
    for k, v in dados.model_dump(exclude={"faixas"}).items():
        setattr(p, k, v)
    p.faixas = [FaixaExcedente(**f.model_dump()) for f in sorted(dados.faixas, key=lambda f: f.de_creditos)]


@router.post("/planos", status_code=201)
def criar_plano(dados: PlanoIn, ctx: EquipeMayaCorp):
    p = Plano()
    _aplicar_plano(p, dados)
    ctx.db.add(p)
    ctx.db.flush()
    ctx.auditar("tribunais.plano.criar", "trb_plano", p.id, descricao=p.nome)
    ctx.db.commit()
    return plano_out(p)


@router.put("/planos/{pid}")
def salvar_plano(pid: int, dados: PlanoIn, ctx: EquipeMayaCorp):
    p = ctx.db.get(Plano, pid)
    if not p:
        raise HTTPException(404, "Plano não encontrado.")
    _aplicar_plano(p, dados)
    ctx.auditar("tribunais.plano.alterar", "trb_plano", p.id, descricao=p.nome)
    ctx.db.commit()
    return plano_out(p)


@router.put("/pesos")
def salvar_pesos(pesos: dict[str, int], ctx: EquipeMayaCorp):
    for tipo, n in pesos.items():
        if tipo not in catalogo.PESOS_PADRAO or n < 0:
            raise HTTPException(400, f"Peso inválido: {tipo}.")
        atual = ctx.db.get(PesoProtocolo, tipo) or PesoProtocolo(tipo_protocolo=tipo, creditos=n)
        atual.creditos = n
        ctx.db.add(atual)
    ctx.db.commit()
    return planos(ctx)


@router.put("/faixas-setup")
def salvar_setup(faixas: list[FaixaSetupIn], ctx: EquipeMayaCorp):
    for f in ctx.db.scalars(select(FaixaSetup)):
        ctx.db.delete(f)
    for f in faixas:
        ctx.db.add(FaixaSetup(**f.model_dump()))
    ctx.db.commit()
    return planos(ctx)


@router.put("/config/plano")
def definir_plano(ctx: EquipeMayaCorp, plano_id: int | None = None):
    cfg = service.config(ctx.db, ctx.tenant_id)
    if plano_id and not ctx.db.get(Plano, plano_id):
        raise HTTPException(404, "Plano não encontrado.")
    cfg.plano_id = plano_id
    ctx.auditar("tribunais.plano.definir", "trb_config", ctx.tenant_id, descricao=str(plano_id))
    ctx.db.commit()
    return _config_out(ctx)


def _consumo_out(c: Consumo) -> dict:
    return {"id": c.id, "ciclo": c.ciclo, "origem": c.origem, "referencia_id": c.referencia_id, "creditos": c.creditos,
            "quantidade": c.quantidade, "dentro_franquia": c.dentro_franquia, "valor": str(c.valor), "motivo": c.motivo,
            "criado_em": _dt(c.criado_em)}


@router.get("/cobranca")
def cobranca_ciclo(ctx: Ctx, ciclo: str | None = Query(None, pattern=r"^\d{4}-\d{2}$")):
    f = cobranca.fechamento(ctx.db, ctx.tenant_id, ciclo or cobranca.ciclo())
    ciclos = sorted(set(ctx.db.scalars(select(Consumo.ciclo).where(Consumo.tenant_id == ctx.tenant_id))) | {cobranca.ciclo()}, reverse=True)
    out = {**{k: (str(v) if isinstance(v, Decimal) else v) for k, v in f.items() if k not in ("lancamentos", "uso")},
           "uso": {k: (str(v) if isinstance(v, Decimal) else v) for k, v in f["uso"].items()},
           "lancamentos": [_consumo_out(c) for c in f["lancamentos"]], "ciclos": ciclos,
           "pesos": {p.tipo_protocolo: p.creditos for p in ctx.db.scalars(select(PesoProtocolo))}}
    ctx.db.commit()
    return out


@router.post("/cobranca/estorno")
def estorno(dados: EstornoIn, ctx: Ctx):
    """Regra 7: estorno manual exige motivo e administrador (fica na auditoria)."""
    if not ctx.usuario.administrador:
        raise HTTPException(403, "Só administradores fazem estorno.")
    c = ctx.db.get(Consumo, dados.consumo_id)
    if not c or c.tenant_id != ctx.tenant_id:
        raise HTTPException(404, "Lançamento não encontrado.")
    try:
        e = cobranca.estornar(ctx.db, c, dados.motivo, ctx.usuario)
    except ValueError as ex:
        raise HTTPException(400, str(ex)) from ex
    ctx.auditar("tribunais.cobranca.estorno", "trb_consumo", e.id, descricao=f"Estorno de #{c.id}: {dados.motivo}")
    ctx.db.commit()
    return _consumo_out(e)


@router.post("/cobranca/ajuste")
def ajuste(dados: AjusteIn, ctx: EquipeMayaCorp):
    ref = (ctx.db.scalar(select(func.max(Consumo.referencia_id)).where(Consumo.tenant_id == ctx.tenant_id, Consumo.origem == "ajuste")) or 0) + 1
    c = Consumo(tenant_id=ctx.tenant_id, ciclo=cobranca.ciclo(), origem="ajuste", referencia_id=ref, creditos=dados.creditos,
                dentro_franquia=dados.valor == 0, valor=dados.valor, motivo=dados.motivo, criado_por_id=ctx.usuario.id)
    ctx.db.add(c)
    ctx.auditar("tribunais.cobranca.ajuste", "trb_consumo", None, descricao=f"{dados.creditos} créditos · {dados.motivo}")
    ctx.db.commit()
    return _consumo_out(c)


# ------------------------------------------------------------------------------------------------ público (recibo)
@publico.get("/recibo")
def recibo_publico(token: str, db: DB):
    """Recibo grande demais para anexar: link do e-mail, válido por 7 dias."""
    try:
        rid = notificacoes.retorno_do_link(token)
    except jwt.PyJWTError:
        raise HTTPException(401, "Link expirado ou inválido. Abra o recibo pelo MayaJus.") from None
    r = db.get(RetornoProtocolo, rid)
    if not r or not r.recibo_ref or not Path(r.recibo_ref).exists():
        raise HTTPException(404, "Recibo não encontrado.")
    return FileResponse(r.recibo_ref, media_type="application/pdf", filename=f"recibo-{r.numero_protocolo or r.id}.pdf")


@router.post("/rotina-diaria")
def rodar_rotina(ctx: EquipeMayaCorp):
    """Dispara na hora a rotina diária (sincronização, certificados, saúde)."""
    jobs.rotina_diaria()
    jobs.disparar()
    return {"ok": True}
